"""Testes: Projetos unificados — área do modelo e o filtro de modelos.

Projetos é um módulo só para PDE, PDR e INOV. O modelo (TipoProjeto) diz de
que área é o projeto, e o filtro `?tipo=`/`?area=` (repetíveis) recorta
listagem, resumo, alertas e export do mesmo jeito.
"""
import io

import pytest


def _novo_tipo(client, h, nome, area):
    res = client.post("/api/tipos-projeto", json={"nome": nome, "area": area}, headers=h)
    assert res.status_code == 201, res.get_json()
    return res.get_json()["tipo"]


def _novo_projeto(client, h, nome, tipo, **extra):
    res = client.post("/api/projetos", json={"nome": nome, "tipo": tipo,
                                             "entregaveis": [], **extra}, headers=h)
    assert res.status_code == 201, res.get_json()
    return res.get_json()["projeto"]["id"]


def _projeto_legado(app, nome="Legado sem modelo", **extra):
    """Projeto de antes dos modelos: `tipo = ""`. Só nasce assim pelo banco."""
    from models import db, Projeto
    with app.app_context():
        p = Projeto(nome=nome, tipo="", ano=2026, status="execucao", **extra)
        db.session.add(p)
        db.session.commit()
        return p.id


def _nomes(client, h, qs=""):
    res = client.get("/api/projetos" + qs, headers=h)
    assert res.status_code == 200, res.get_json()
    return sorted(p["nome"] for p in res.get_json()["projetos"])


@pytest.fixture
def portfolio(app, client, admin_token, auth_headers):
    """Um projeto por área, mais um OEM (tipo antigo, área pde) e um legado."""
    h = auth_headers(admin_token)
    _novo_tipo(client, h, "PDE RUO", "pde")
    _novo_tipo(client, h, "PDR RUO", "pdr")
    _novo_tipo(client, h, "PDR IVD", "pdr")
    _novo_tipo(client, h, "INOV", "inov")
    _novo_projeto(client, h, "Leitor", "PDE RUO")
    _novo_projeto(client, h, "Kit A", "PDR RUO")
    _novo_projeto(client, h, "Kit B", "PDR IVD")
    _novo_projeto(client, h, "Ideia", "INOV")
    _novo_projeto(client, h, "Centrífuga", "OEM")
    _projeto_legado(app)
    return h


# ── Área do modelo ───────────────────────────────────────────────────────────

def test_tipos_trazem_area_e_as_areas_da_tela(client, admin_token, auth_headers):
    r = client.get("/api/tipos-projeto", headers=auth_headers(admin_token)).get_json()
    oem = next(t for t in r["tipos"] if t["nome"] == "OEM")
    assert oem["area"] == "pde"           # tipos anteriores à unificação são do PDE
    assert [a["slug"] for a in r["areas"]] == ["pde", "pdr", "inov"]
    assert {a["sigla"] for a in r["areas"]} == {"PDE", "PDR", "INOV"}
    assert all(a["accent"].startswith("#") for a in r["areas"])
    assert r["area_legado"] == "pde"


def test_criar_tipo_exige_area_valida(client, admin_token, auth_headers):
    h = auth_headers(admin_token)
    assert client.post("/api/tipos-projeto", json={"nome": "Sem área"},
                       headers=h).status_code == 400
    assert client.post("/api/tipos-projeto", json={"nome": "Área torta", "area": "xyz"},
                       headers=h).status_code == 400
    t = _novo_tipo(client, h, "PDR IVD", "PDR")    # slug sem diferenciar maiúsculas
    assert t["area"] == "pdr"


def test_trocar_area_do_tipo_move_os_projetos_dele(client, admin_token, auth_headers):
    h = auth_headers(admin_token)
    t = _novo_tipo(client, h, "Kits", "pde")
    _novo_projeto(client, h, "Kit X", "Kits")
    assert _nomes(client, h, "?area=pdr") == []

    res = client.put(f"/api/tipos-projeto/{t['id']}", json={"area": "pdr"}, headers=h)
    assert res.status_code == 200, res.get_json()
    assert res.get_json()["tipo"]["area"] == "pdr"
    assert res.get_json()["tipo"]["nome"] == "Kits"      # só a área mudou
    # a área do projeto vem do modelo: nada a propagar
    assert _nomes(client, h, "?area=pdr") == ["Kit X"]


def test_editar_tipo_valida_area_e_exige_algum_campo(client, admin_token, auth_headers):
    h = auth_headers(admin_token)
    t = _novo_tipo(client, h, "Kits", "pde")
    assert client.put(f"/api/tipos-projeto/{t['id']}", json={"area": "xyz"},
                      headers=h).status_code == 400
    assert client.put(f"/api/tipos-projeto/{t['id']}", json={},
                      headers=h).status_code == 400
    # renomear e trocar a área na mesma chamada
    res = client.put(f"/api/tipos-projeto/{t['id']}",
                     json={"nome": "INOV Kits", "area": "inov"}, headers=h).get_json()
    assert (res["tipo"]["nome"], res["tipo"]["area"]) == ("INOV Kits", "inov")


def test_gestor_cria_modelo_de_qualquer_area(client, gestor_token, auth_headers):
    """Nesta etapa não há distinção de usuário por área: o gestor cria em todas."""
    t = _novo_tipo(client, auth_headers(gestor_token), "INOV", "inov")
    assert t["area"] == "inov"


# ── Modelo obrigatório no projeto ────────────────────────────────────────────

def test_projeto_novo_exige_modelo(client, admin_token, auth_headers):
    res = client.post("/api/projetos", json={"nome": "Sem modelo"},
                      headers=auth_headers(admin_token))
    assert res.status_code == 400
    assert "modelo" in res.get_json()["erro"]


def test_projeto_com_modelo_nao_volta_a_ficar_sem(client, admin_token, auth_headers):
    h = auth_headers(admin_token)
    pid = _novo_projeto(client, h, "P", "OEM")
    res = client.put(f"/api/projetos/{pid}", json={"tipo": ""}, headers=h)
    assert res.status_code == 400


def test_projeto_legado_sem_modelo_continua_editavel(app, client, admin_token, auth_headers):
    """O formulário reenvia `tipo: ""` do legado; isso não pode travar a edição."""
    h = auth_headers(admin_token)
    pid = _projeto_legado(app)
    res = client.put(f"/api/projetos/{pid}", json={"tipo": "", "nome": "Renomeado"},
                     headers=h)
    assert res.status_code == 200, res.get_json()
    # e dá para classificá-lo depois
    res = client.put(f"/api/projetos/{pid}", json={"tipo": "OEM"}, headers=h)
    assert res.status_code == 200
    assert res.get_json()["projeto"]["tipo"] == "OEM"


# ── Filtro de modelos (?tipo= / ?area=) ──────────────────────────────────────

def test_sem_filtro_lista_todas_as_areas(client, portfolio):
    assert _nomes(client, portfolio) == sorted(
        ["Leitor", "Kit A", "Kit B", "Ideia", "Centrífuga", "Legado sem modelo"])


def test_filtro_por_varios_modelos(client, portfolio):
    assert _nomes(client, portfolio, "?tipo=PDE RUO&tipo=PDR IVD") == ["Kit B", "Leitor"]
    # sem diferenciar maiúsculas, como o cadastro
    assert _nomes(client, portfolio, "?tipo=pdr ruo") == ["Kit A"]


def test_filtro_por_area_inclui_todos_os_modelos_dela(client, portfolio):
    assert _nomes(client, portfolio, "?area=pdr") == ["Kit A", "Kit B"]
    assert _nomes(client, portfolio, "?area=inov") == ["Ideia"]


def test_area_pde_inclui_tipos_antigos_e_projeto_sem_modelo(client, portfolio):
    assert _nomes(client, portfolio, "?area=pde") == sorted(
        ["Leitor", "Centrífuga", "Legado sem modelo"])


def test_area_e_modelo_combinam_por_uniao(client, portfolio):
    """Marcar a área PDR inteira e só o INOV traz os dois recortes."""
    assert _nomes(client, portfolio, "?area=pdr&tipo=INOV") == ["Ideia", "Kit A", "Kit B"]


def test_area_sem_modelo_cadastrado_devolve_vazio(client, admin_token, auth_headers):
    """Entrar pelo hub do INOV antes de existir modelo INOV não pode mostrar
    o portfólio inteiro."""
    h = auth_headers(admin_token)
    _novo_projeto(client, h, "Centrífuga", "OEM")
    assert _nomes(client, h, "?area=inov") == []


def test_area_invalida_retorna_400(client, portfolio):
    assert client.get("/api/projetos?area=xyz", headers=portfolio).status_code == 400


def test_resumo_respeita_o_filtro(client, portfolio):
    r = client.get("/api/entregaveis/resumo?area=pdr", headers=portfolio).get_json()
    assert r["projetos"] == 2
    r = client.get("/api/entregaveis/resumo", headers=portfolio).get_json()
    assert r["projetos"] == 6


def test_alertas_respeitam_o_filtro(client, admin_token, auth_headers):
    h = auth_headers(admin_token)
    _novo_tipo(client, h, "PDR RUO", "pdr")
    vencido = {"data_inicio_prev": "2020-01-01", "data_fim_prev": "2020-12-31"}
    pde = _novo_projeto(client, h, "Venceu PDE", "OEM", **vencido)
    pdr = _novo_projeto(client, h, "Venceu PDR", "PDR RUO", **vencido)

    def vencidos(qs=""):
        r = client.get("/api/projetos/alertas" + qs, headers=h).get_json()
        return {a["projeto_id"] for a in r["alertas"] if a["tipo"] == "projeto_vencido"}

    assert vencidos() == {pde, pdr}
    assert vencidos("?area=pdr") == {pdr}
    assert vencidos("?tipo=OEM") == {pde}


def test_export_filtra_e_traz_modelo_e_area(client, portfolio):
    from openpyxl import load_workbook
    res = client.get("/api/entregaveis/export?area=pdr", headers=portfolio)
    assert res.status_code == 200
    wb = load_workbook(io.BytesIO(res.data))

    ws = wb.worksheets[0]
    assert [c.value for c in ws[1][:3]] == ["Projeto", "Modelo", "Área"]
    linhas = sorted((r[0], r[1], r[2]) for r in ws.iter_rows(min_row=2, values_only=True))
    assert linhas == [("Kit A", "PDR RUO", "PDR"), ("Kit B", "PDR IVD", "PDR")]

    pmo = wb["PMO"]
    assert [c.value for c in pmo[1][:4]] == ["Projeto", "Status", "Modelo", "Área"]


def test_export_mostra_legado_como_pde(client, portfolio):
    from openpyxl import load_workbook
    res = client.get("/api/entregaveis/export?area=pde", headers=portfolio)
    ws = load_workbook(io.BytesIO(res.data)).worksheets[0]
    areas = {r[0]: r[2] for r in ws.iter_rows(min_row=2, values_only=True)}
    assert areas == {"Leitor": "PDE", "Centrífuga": "PDE", "Legado sem modelo": "PDE"}
