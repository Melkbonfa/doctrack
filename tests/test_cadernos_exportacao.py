"""Cadernos: anexo dentro da nota e exportação em zip para a liderança.

O que estes testes fixam, e por quê:

  * **Anexo no texto.** O link que o editor põe para um arquivo da página é
    guardado como `<a data-arquivo>` — sem URL e sem o token da sessão —, o
    painel marca o arquivo como "no texto" e o link baixa a versão VIGENTE.
  * **O zip leva o que foi revisto.** A prévia marca por padrão o "Exporta" de
    cada arquivo e tudo que está citado no texto; o zip leva só o que veio
    marcado, com o relatório .docx, as figuras originais e o LEIA-ME com o
    SHA-256 de cada arquivo.
  * **O .docx preserva o que importa num registro de teste**: cor de célula,
    células mescladas, realce, listas e o link para o arquivo dentro do zip.
"""
import hashlib
import io
import re
import zipfile

import pytest
from docx import Document
from PIL import Image

import arquivos_store
from models import db, User
from cadernos import exportacao, relatorio_docx
from cadernos.models import CadernoExportacao, CadernoPagina


@pytest.fixture(autouse=True)
def _pastas_temporarias(tmp_path, monkeypatch):
    monkeypatch.setattr(arquivos_store, "RAIZ", str(tmp_path / "arquivos"))
    monkeypatch.setattr(exportacao, "PASTA_TEMP", str(tmp_path / "exportacoes"))
    yield


@pytest.fixture
def area_pdr(app):
    for email in ("gestor@test.com", "tecnico@test.com", "leitura@test.com"):
        User.query.filter_by(email=email).first().areas = "pdr"
    db.session.commit()


def H(token):
    return {"Authorization": f"Bearer {token}"}


def _png(largura=40, altura=20):
    buf = io.BytesIO()
    Image.new("RGB", (largura, altura), (200, 30, 30)).save(buf, "PNG")
    return buf.getvalue()


def _caderno_com_pagina(client, token):
    res = client.post("/cadernos/api/cadernos", json={"area": "pdr", "nome": "Projeto L-000",
                                                      "modelo": "projeto_pdr"}, headers=H(token))
    cad = res.get_json()["caderno"]["uuid"]
    arv = client.get(f"/cadernos/api/cadernos/{cad}/arvore", headers=H(token)).get_json()
    secao = next(n["uuid"] for n in arv["nos"] if n["nome"] == "Desenvolvimento")
    p = client.post(f"/cadernos/api/secoes/{secao}/paginas", json={"tipo": "registro_teste"},
                    headers=H(token)).get_json()["pagina"]
    return cad, secao, p


def _envia(client, token, cad, escopo, ref, nome, conteudo=b"dados"):
    res = client.post("/cadernos/api/arquivos", data={
        "caderno": cad, "escopo": escopo, "ref": ref, "arquivo": (io.BytesIO(conteudo), nome)},
        content_type="multipart/form-data", headers=H(token))
    assert res.status_code == 201, res.get_json()
    return res.get_json()["arquivo"]


def _grava(client, token, pagina, html, titulo=None, ficha=None):
    url = f"/cadernos/api/paginas/{pagina}"
    client.post(url + "/edicao", headers=H(token))
    versao = client.get(url, headers=H(token)).get_json()["pagina"]["versao"]
    corpo = {"versao_base": versao, "html": html}
    if titulo:
        corpo["titulo"] = titulo
    if ficha:
        corpo["ficha"] = ficha
    res = client.put(url, json=corpo, headers=H(token))
    assert res.status_code == 200, res.get_json()
    client.delete(url + "/edicao", headers=H(token))


# ══════════════════════════════════════════════════════════════════════════════
# ANEXO NO TEXTO
# ══════════════════════════════════════════════════════════════════════════════

def test_anexo_no_texto_e_guardado_sem_url_nem_token(client, tecnico_token, area_pdr):
    cad, _, p = _caderno_com_pagina(client, tecnico_token)
    a = _envia(client, tecnico_token, cad, "pagina", p["uuid"], "corrida_0912.tlpd")
    href = f"/cadernos/api/arquivos/{a['uuid']}/conteudo?token={tecnico_token}&download=1&atual=1"
    _grava(client, tecnico_token, p["uuid"], f'<p>Corrida: <a href="{href}">corrida_0912.tlpd</a></p>')
    html = CadernoPagina.query.filter_by(uuid=p["uuid"]).first().conteudo_html
    assert f'<a data-arquivo="{a["uuid"]}">corrida_0912.tlpd</a>' in html
    assert tecnico_token not in html and "href" not in html

    lista = client.get(f"/cadernos/api/arquivos?caderno={cad}&escopo=pagina&ref={p['uuid']}",
                       headers=H(tecnico_token)).get_json()["arquivos"]
    assert [x["no_texto"] for x in lista] == [True]


def test_anexo_no_texto_baixa_a_versao_vigente(client, tecnico_token, area_pdr):
    cad, _, p = _caderno_com_pagina(client, tecnico_token)
    v1 = _envia(client, tecnico_token, cad, "pagina", p["uuid"], "analise.xlsx", b"v1")
    client.post(f"/cadernos/api/arquivos/{v1['uuid']}/versao", data={
        "arquivo": (io.BytesIO(b"v2"), "analise.xlsx")}, content_type="multipart/form-data",
        headers=H(tecnico_token))
    url = f"/cadernos/api/arquivos/{v1['uuid']}/conteudo"
    assert client.get(url, headers=H(tecnico_token)).data == b"v1"               # versão pedida
    assert client.get(url + "?atual=1", headers=H(tecnico_token)).data == b"v2"  # a do link no texto
    # O link continua marcando a família como "no texto" depois da versão nova.
    _grava(client, tecnico_token, p["uuid"], f'<p><a data-arquivo="{v1["uuid"]}">analise.xlsx</a></p>')
    lista = client.get(f"/cadernos/api/arquivos?caderno={cad}&escopo=pagina&ref={p['uuid']}",
                       headers=H(tecnico_token)).get_json()["arquivos"]
    assert [(x["versao"], x["no_texto"]) for x in lista] == [(2, True)]


# ══════════════════════════════════════════════════════════════════════════════
# EXPORTAÇÃO
# ══════════════════════════════════════════════════════════════════════════════

def _cenario(client, token):
    """Página com imagem, anexo no texto (com Exporta DESLIGADO), arquivo de página
    ligado, arquivo da seção desligado e programa do caderno desligado."""
    cad, secao, p = _caderno_com_pagina(client, token)
    img = client.post(f"/cadernos/api/paginas/{p['uuid']}/imagens",
                      data={"upload": (io.BytesIO(_png()), "curva.png")},
                      content_type="multipart/form-data", headers=H(token)).get_json()["uuid"]
    citado = _envia(client, token, cad, "pagina", p["uuid"], "corrida.tlpd", b"corrida")
    client.patch(f"/cadernos/api/arquivos/{citado['uuid']}", json={"incluir_exportacao": False},
                 headers=H(token))
    laudo = _envia(client, token, cad, "pagina", p["uuid"], "laudo.pdf", b"%PDF-1.4 laudo")
    da_secao = _envia(client, token, cad, "secao", secao, "planilha_geral.xlsx", b"xlsx")
    programa = _envia(client, token, cad, "caderno", "", "analise_placa.exe", b"MZ programa")
    html = (f'<h2>1. DESCRIÇÃO</h2><p>Arquivo: <a data-arquivo="{citado["uuid"]}">corrida.tlpd</a></p>'
            f'<figure class="image"><img data-arquivo="{img}"><figcaption>Curva padrão</figcaption></figure>'
            '<h2>4. CONCLUSÃO</h2><p>Curva aprovada com R² 0,99.</p>')
    _grava(client, token, p["uuid"], html, titulo="Teste 56 – Curva",
           ficha={"projeto": "L-000", "numero": "56", "inicio": "2026-03-20", "termino": "2026-03-21",
                  "resultado": "aprovado"})
    return cad, p, {"citado": citado, "laudo": laudo, "secao": da_secao, "programa": programa}


def test_previa_marca_exporta_e_o_que_esta_no_texto(client, tecnico_token, area_pdr):
    cad, p, arq = _cenario(client, tecnico_token)
    res = client.post(f"/cadernos/api/cadernos/{cad}/exportacao/previa", json={"paginas": [p["uuid"]]},
                      headers=H(tecnico_token))
    assert res.status_code == 200, res.get_json()
    r = res.get_json()
    assert r["escopo"] == "Página Teste 56 – Curva" and r["imagens"] == 1
    marcado = {a["nome"]: a["marcado"] for a in r["arquivos"]}
    assert marcado == {"corrida.tlpd": True,          # Exporta desligado, mas está no texto
                       "laudo.pdf": True,             # Exporta ligado (padrão da página)
                       "planilha_geral.xlsx": False,  # seção: desligado por padrão
                       "analise_placa.exe": False}    # caderno: desligado por padrão
    destinos = {a["nome"]: a["destino"] for a in r["arquivos"]}
    assert destinos["corrida.tlpd"] == "registros/01 - Teste 56 – Curva/corrida.tlpd"
    assert destinos["planilha_geral.xlsx"] == "projeto/Desenvolvimento/planilha_geral.xlsx"
    assert destinos["analise_placa.exe"] == "projeto/analise_placa.exe"


def test_zip_leva_o_relatorio_as_figuras_e_so_os_arquivos_marcados(client, tecnico_token, area_pdr):
    cad, p, arq = _cenario(client, tecnico_token)
    escolhidos = [arq["citado"]["uuid"], arq["laudo"]["uuid"], arq["programa"]["uuid"]]
    res = client.post(f"/cadernos/api/cadernos/{cad}/exportacoes",
                      json={"paginas": [p["uuid"]], "arquivos": escolhidos}, headers=H(tecnico_token))
    assert res.status_code == 202, res.get_json()
    t = res.get_json()["tarefa"]
    assert t["estado"] == "pronta", t                 # em teste a tarefa roda na hora
    exp = t["resultado"]["exportacao"]
    assert client.get(f"/cadernos/api/tarefas/{t['id']}", headers=H(tecnico_token)).get_json()["tarefa"]["estado"] == "pronta"

    baixa = client.get(f"/cadernos/api/exportacoes/{exp['uuid']}/zip?token={tecnico_token}")
    assert baixa.status_code == 200 and baixa.headers["Content-Type"] == "application/zip"
    zf = zipfile.ZipFile(io.BytesIO(baixa.data))
    nomes = set(zf.namelist())
    assert {"LEIA-ME.txt", "Relatorio.docx",
            "registros/01 - Teste 56 – Curva/corrida.tlpd",
            "registros/01 - Teste 56 – Curva/laudo.pdf",
            "projeto/analise_placa.exe"} <= nomes
    assert "projeto/Desenvolvimento/planilha_geral.xlsx" not in nomes          # desmarcado
    assert [n for n in nomes if n.startswith("imagens/")] == ["imagens/Figura 001 - Curva padrão.png"]

    leia = zf.read("LEIA-ME.txt").decode("utf-8-sig")
    assert hashlib.sha256(b"corrida").hexdigest() in leia
    assert hashlib.sha256(zf.read("Relatorio.docx")).hexdigest() in leia

    doc = Document(io.BytesIO(zf.read("Relatorio.docx")))
    texto = "\n".join(par.text for par in doc.paragraphs)
    assert "Teste 56 – Curva" in texto and "Figura 1 — Curva padrão" in texto
    resumo = doc.tables[1]                            # capa, testes efetuados, ...
    linha = [c.text for c in resumo.rows[1].cells]
    assert linha[0] == "56" and linha[4] == "Aprovado" and "R² 0,99" in linha[5]
    # O link do texto aponta para o arquivo DENTRO do zip (caminho relativo).
    alvos = [r.target_ref for r in doc.part.rels.values() if r.reltype.endswith("/hyperlink")]
    assert any(a.endswith("corrida.tlpd") and a.startswith("registros/") for a in alvos)

    e = CadernoExportacao.query.filter_by(uuid=exp["uuid"]).first()
    assert '"versao"' in e.paginas_json and hashlib.sha256(b"corrida").hexdigest() in e.manifesto_json
    hist = client.get(f"/cadernos/api/cadernos/{cad}/exportacoes", headers=H(tecnico_token)).get_json()
    assert hist["exportacoes"][0]["uuid"] == exp["uuid"] and hist["exportacoes"][0]["disponivel"]


def test_arquivo_desmarcado_vira_texto_no_relatorio(client, tecnico_token, area_pdr):
    cad, p, arq = _cenario(client, tecnico_token)
    t = client.post(f"/cadernos/api/cadernos/{cad}/exportacoes",
                    json={"paginas": [p["uuid"]], "arquivos": []}, headers=H(tecnico_token)).get_json()["tarefa"]
    zf = zipfile.ZipFile(exportacao.caminho_zip(t["resultado"]["exportacao"]["uuid"]))
    doc = Document(io.BytesIO(zf.read("Relatorio.docx")))
    assert "corrida.tlpd (arquivo fora do pacote)" in "\n".join(par.text for par in doc.paragraphs)


def test_so_editor_exporta_e_quem_nao_ve_o_caderno_nem_sabe(client, tecnico_token, leitura_token,
                                                           admin_token, area_pdr):
    """R-50: exportar é saída de dados, e o perfil leitura só consulta. Antes
    bastava ler o caderno (PAPEL_EXPORTAR = leitor)."""
    from models import AuditLog
    cad, p, _ = _cenario(client, tecnico_token)
    corpo = {"paginas": [p["uuid"]], "arquivos": []}

    # leitura lê o caderno, mas não tira prévia, não gera e não baixa
    assert client.get(f"/cadernos/api/cadernos/{cad}", headers=H(leitura_token)).status_code == 200
    assert client.post(f"/cadernos/api/cadernos/{cad}/exportacao/previa", json=corpo,
                       headers=H(leitura_token)).status_code == 403
    assert client.post(f"/cadernos/api/cadernos/{cad}/exportacoes", json=corpo,
                       headers=H(leitura_token)).status_code == 403
    assert CadernoExportacao.query.count() == 0

    # o editor (técnico dono) exporta; o zip dele também não desce para o leitura
    t = client.post(f"/cadernos/api/cadernos/{cad}/exportacoes", json=corpo,
                    headers=H(tecnico_token))
    assert t.status_code == 202
    uuid_ = t.get_json()["tarefa"]["resultado"]["exportacao"]["uuid"]
    assert client.get(f"/cadernos/api/exportacoes/{uuid_}/zip",
                      headers=H(leitura_token)).status_code == 403
    assert client.get(f"/cadernos/api/exportacoes/{uuid_}/zip",
                      headers=H(tecnico_token)).status_code == 200

    # o histórico do que já saiu é registro, não saída: quem lê o caderno vê
    hist = client.get(f"/cadernos/api/cadernos/{cad}/exportacoes", headers=H(leitura_token))
    assert hist.status_code == 200 and len(hist.get_json()["exportacoes"]) == 1

    negados = AuditLog.query.filter_by(acao="ACESSO_NEGADO", usuario_email="leitura@test.com").all()
    assert {l.valor_novo for l in negados} == {"papel leitor no caderno Projeto L-000 (exige editor)"}

    # Sem a área PDR, o caderno nem existe para a pessoa (404 não revela).
    User.query.filter_by(email="leitura@test.com").first().areas = ""
    db.session.commit()
    assert client.post(f"/cadernos/api/cadernos/{cad}/exportacao/previa", json=corpo,
                       headers=H(leitura_token)).status_code == 404


def test_tarefa_so_para_quem_iniciou(client, tecnico_token, gestor_token, area_pdr):
    cad, p, _ = _cenario(client, tecnico_token)
    t = client.post(f"/cadernos/api/cadernos/{cad}/exportacoes", json={"paginas": [p["uuid"]], "arquivos": []},
                    headers=H(tecnico_token)).get_json()["tarefa"]
    assert client.get(f"/cadernos/api/tarefas/{t['id']}", headers=H(gestor_token)).status_code == 404


def test_zip_expirado_responde_410(client, tecnico_token, area_pdr):
    cad, p, _ = _cenario(client, tecnico_token)
    t = client.post(f"/cadernos/api/cadernos/{cad}/exportacoes", json={"paginas": [p["uuid"]], "arquivos": []},
                    headers=H(tecnico_token)).get_json()["tarefa"]
    uuid_ = t["resultado"]["exportacao"]["uuid"]
    import os
    os.remove(exportacao.caminho_zip(uuid_))
    assert client.get(f"/cadernos/api/exportacoes/{uuid_}/zip", headers=H(tecnico_token)).status_code == 410


def test_sem_paginas_nao_exporta(client, tecnico_token, area_pdr):
    cad, _, _ = _caderno_com_pagina(client, tecnico_token)
    assert client.post(f"/cadernos/api/cadernos/{cad}/exportacao/previa", json={"paginas": []},
                       headers=H(tecnico_token)).status_code == 400


def test_nome_seguro_para_o_windows():
    assert exportacao.nome_seguro('Teste 5: lote "A"/B?') == "Teste 5 lote A B"
    assert exportacao.nome_seguro("CON") == "_CON"
    assert exportacao.nome_seguro("...") == "sem nome"
    assert len(exportacao.nome_seguro("x" * 200, 80)) == 80


# ══════════════════════════════════════════════════════════════════════════════
# HTML → WORD
# ══════════════════════════════════════════════════════════════════════════════

def _docx(html, imagens=None, links=None):
    rel = relatorio_docx.Relatorio(caderno="C", escopo="E", emitido_por="Fulana",
                                   imagem_de=lambda u: (imagens or {}).get(u),
                                   link_de=lambda u: (links or {}).get(u))
    pg = relatorio_docx.PaginaRelatorio(titulo="P", secao="S", tipo="nota", html=html, versao=1)
    dados = rel.gerar([pg], [])
    return Document(io.BytesIO(dados)), rel


def test_docx_tabela_com_cor_mescla_e_cabecalho():
    doc, _ = _docx('<figure class="table"><table><colgroup><col style="width:25%"><col style="width:75%"></colgroup>'
                   '<thead><tr><th>Reagente</th><th>Lote</th></tr></thead>'
                   '<tbody><tr><td rowspan="2" style="background-color:#9BBB59">Mix</td><td>2410</td></tr>'
                   '<tr><td>2411</td></tr><tr><td colspan="2">Total</td></tr></tbody></table></figure>')
    t = next(t for t in doc.tables if t.rows[0].cells[0].text == "Reagente")
    xml = t._tbl.xml
    assert 'w:fill="9BBB59"' in xml
    assert 'w:vMerge' in xml and 'w:gridSpan w:val="2"' in xml
    assert "w:tblHeader" in xml
    assert [c.text for c in t.rows[0].cells] == ["Reagente", "Lote"]


def test_docx_realce_lista_numerada_e_link_externo():
    doc, _ = _docx('<p><mark class="marker-yellow">atenção</mark> e Cy<sup>5</sup></p>'
                   '<ol><li>Selar a placa</li><li>Centrifugar<ul><li>1 min</li></ul></li></ol>'
                   '<p><a href="https://loccus.sharepoint.com/x">corrida</a></p>')
    corpo = doc.element.body.xml
    assert 'w:highlight w:val="yellow"' in corpo
    assert 'w:vertAlign w:val="superscript"' in corpo
    textos = [p.text for p in doc.paragraphs]
    assert "1.\tSelar a placa" in textos and "2.\tCentrifugar" in textos and "◦\t1 min" in textos   # sublista: marcador do 2º nível
    assert any(r.target_ref == "https://loccus.sharepoint.com/x" for r in doc.part.rels.values())


def test_docx_imagem_larga_e_reduzida_e_numerada():
    png = _png(3000, 600)
    doc, rel = _docx('<figure class="image"><img data-arquivo="u1"><figcaption>Gel</figcaption></figure>'
                     '<figure class="image"><img data-arquivo="sumiu"></figure>',
                     imagens={"u1": (png, ".png")})
    assert [(f.numero, f.uuid, f.legenda) for f in rel.figuras] == [(1, "u1", "Gel")]
    blob = next(p for p in doc.part.package.parts if p.partname.startswith("/word/media/")).blob
    assert Image.open(io.BytesIO(blob)).size[0] < 3000           # reamostrada para o Word
    textos = [p.text for p in doc.paragraphs]
    assert "Figura 1 — Gel" in textos and "[imagem indisponível]" in textos


def test_conclusao_vem_da_secao_conclusao_do_modelo():
    html = ('<h2>3. RESULTADOS</h2><p>ct 24</p><h2>4. CONCLUSÃO</h2><p>Aprovado.</p><h3>4.1 Notas</h3>'
            '<p>sem desvio</p><h2>5. ANEXOS</h2><p>x</p>')
    assert relatorio_docx.conclusao_de(html) == "Aprovado. 4.1 Notas sem desvio"
    assert relatorio_docx.conclusao_de("<p>sem seção</p>") == ""
    assert relatorio_docx.cor_hex("rgb(155, 187, 89)") == "9BBB59"
    assert relatorio_docx.cor_hex("#abc") == "AABBCC"
