"""
cadernos/relatorio_docx.py — Páginas do caderno → relatório .docx da exportação.

Monta o `Relatorio.docx` que vai no zip para a liderança (PLANO_cadernos.md 2.7):
capa, tabela dos testes efetuados, cada página com a ficha e o conteúdo, a lista
de arquivos do pacote e o campo de assinaturas.

**O HTML que chega aqui já passou por `core.sanear_html`.** É um conjunto pequeno
e fechado de tags (`core._TAGS`), sempre bem aninhado — por isso a conversão é
uma árvore simples com uma regra por tag, e não um conversor genérico de HTML.
Tag que não estiver aqui vira o texto dela.

**Modelo do Word.** Se existir `cadernos/modelos/relatorio.docx`, o relatório
nasce dele: estilos, cabeçalho, rodapé e logo vêm de lá, e é assim que o PDR
ajusta a cara do relatório sem mexer em código. O corpo do modelo é descartado.
Sem o arquivo, vale o modelo padrão do python-docx com as margens e fontes
ajustadas aqui.

**Imagens.** Entram na largura útil da página (ou da célula), reamostradas para
no máximo ~200 dpi na largura impressa; o original vai inteiro na pasta
`imagens/` do zip. Toda imagem em bloco ganha "Figura N"; a numeração é a mesma
dos arquivos de `imagens/` (o chamador recebe a lista em `Relatorio.figuras`).
"""
import io
import os
import re
import unicodedata
from dataclasses import dataclass, field
from datetime import datetime
from html.parser import HTMLParser
from urllib.parse import quote

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

MODELO = os.path.join(os.path.dirname(os.path.abspath(__file__)), "modelos", "relatorio.docx")

DPI_MAX = 200                 # resolução máxima da imagem na largura impressa
FONTE_TABELA = Pt(9)          # tabela de placa com 13 colunas precisa caber
CINZA_TH = "F3F4F6"
AZUL_LINK = RGBColor(0x05, 0x63, 0xC1)

REALCES = {
    # As seis cores de realce do OneNote (as mesmas do editor) → realce do Word.
    "marker-yellow": WD_COLOR_INDEX.YELLOW, "marker-lime": WD_COLOR_INDEX.BRIGHT_GREEN,
    "marker-fuchsia": WD_COLOR_INDEX.PINK, "marker-red": WD_COLOR_INDEX.RED,
    "marker-aqua": WD_COLOR_INDEX.TURQUOISE, "marker-silver": WD_COLOR_INDEX.GRAY_25,
    # Canetas do CKEditor: texto colorido, sem realce.
}
CANETAS = {"pen-red": "E71313", "pen-green": "128A00"}
TAMANHOS = {"text-tiny": 0.7, "text-small": 0.85, "text-big": 1.4, "text-huge": 1.8}
CORES_NOMEADAS = {"black": "000000", "white": "FFFFFF", "red": "FF0000", "green": "008000",
                  "blue": "0000FF", "yellow": "FFFF00", "gray": "808080", "grey": "808080",
                  "silver": "C0C0C0", "orange": "FFA500", "purple": "800080"}
RESULTADO_ROTULO = {"": "—", "em_andamento": "Em andamento", "aprovado": "Aprovado",
                    "reprovado": "Reprovado", "inconclusivo": "Inconclusivo"}

_BLOCOS = {"p", "h1", "h2", "h3", "h4", "ul", "ol", "li", "table", "figure", "figcaption",
           "blockquote", "pre", "hr", "div", "thead", "tbody", "tfoot", "tr", "colgroup"}
_VAZIAS = {"img", "br", "hr", "col", "input"}


# ══════════════════════════════════════════════════════════════════════════════
# ÁRVORE DO HTML SANEADO
# ══════════════════════════════════════════════════════════════════════════════

class No:
    __slots__ = ("tag", "attrs", "filhos")

    def __init__(self, tag, attrs=None):
        self.tag = tag
        self.attrs = attrs or {}
        self.filhos = []

    def classes(self):
        return set((self.attrs.get("class") or "").split())

    def estilo(self):
        d = {}
        for decl in (self.attrs.get("style") or "").split(";"):
            if ":" in decl:
                k, v = decl.split(":", 1)
                d[k.strip().lower()] = v.strip()
        return d

    def texto(self):
        partes = []
        for f in self.filhos:
            if isinstance(f, str):
                partes.append(f)
            elif f.tag == "br":
                partes.append("\n")
            else:
                partes.append(f.texto())
                if f.tag in _BLOCOS:
                    partes.append(" ")
        return "".join(partes)


class _Arvore(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.raiz = No("raiz")
        self.pilha = [self.raiz]

    def handle_starttag(self, tag, attrs):
        no = No(tag, {k: (v or "") for k, v in attrs})
        self.pilha[-1].filhos.append(no)
        if tag not in _VAZIAS:
            self.pilha.append(no)

    def handle_startendtag(self, tag, attrs):
        self.pilha[-1].filhos.append(No(tag, {k: (v or "") for k, v in attrs}))

    def handle_endtag(self, tag):
        for i in range(len(self.pilha) - 1, 0, -1):
            if self.pilha[i].tag == tag:
                del self.pilha[i:]
                return

    def handle_data(self, data):
        self.pilha[-1].filhos.append(data)


def arvore(html_texto):
    p = _Arvore()
    p.feed(html_texto or "")
    p.close()
    return p.raiz


# ══════════════════════════════════════════════════════════════════════════════
# UTILIDADES DE ESTILO E XML
# ══════════════════════════════════════════════════════════════════════════════

def _sem_acento(s):
    return unicodedata.normalize("NFKD", s or "").encode("ASCII", "ignore").decode("ASCII").lower()


def cor_hex(valor):
    """'#9BBB59', '#9bb', 'rgb(155,187,89)', 'red' → '9BBB59'. None se não der."""
    v = (valor or "").strip().lower()
    if not v:
        return None
    if v in CORES_NOMEADAS:
        return CORES_NOMEADAS[v]
    m = re.fullmatch(r"#([0-9a-f]{6})", v)
    if m:
        return m.group(1).upper()
    m = re.fullmatch(r"#([0-9a-f])([0-9a-f])([0-9a-f])", v)
    if m:
        return "".join(c * 2 for c in m.groups()).upper()
    m = re.fullmatch(r"rgba?\(\s*(\d+)\s*,\s*(\d+)\s*,\s*(\d+)[^)]*\)", v)
    if m:
        return "".join(f"{min(255, int(x)):02X}" for x in m.groups())
    return None


def _cm_de_css(valor, referencia_cm):
    """'400px' → cm; '53.4%' → % da referência; '12pt' → cm. None se não der."""
    v = (valor or "").strip().lower()
    m = re.fullmatch(r"([\d.]+)\s*(px|%|pt|cm)?", v)
    if not m:
        return None
    n, unidade = float(m.group(1)), m.group(2) or "px"
    if unidade == "px":
        return n * 2.54 / 96
    if unidade == "pt":
        return n * 2.54 / 72
    if unidade == "cm":
        return n
    return referencia_cm * n / 100


def _sombrear_celula(cell, hex6):
    tcPr = cell._tc.get_or_add_tcPr()
    for velho in tcPr.findall(qn("w:shd")):
        tcPr.remove(velho)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex6)
    tcPr.insert_element_before(shd, "w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText",
                               "w:vAlign", "w:hideMark", "w:headers", "w:cellIns", "w:cellDel",
                               "w:cellMerge", "w:tcPrChange")


def _sombrear_run(run, hex6):
    rPr = run._r.get_or_add_rPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex6)
    rPr.insert_element_before(shd, "w:fitText", "w:vertAlign", "w:rtl", "w:cs", "w:em", "w:lang",
                              "w:eastAsianLayout", "w:specVanish", "w:oMath")


def _bordas_tabela(tabela, cor="BFBFBF"):
    tblPr = tabela._tbl.tblPr
    bordas = OxmlElement("w:tblBorders")
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
        b = OxmlElement(f"w:{lado}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), "4")
        b.set(qn("w:space"), "0")
        b.set(qn("w:color"), cor)
        bordas.append(b)
    tblPr.insert_element_before(bordas, "w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook",
                                "w:tblCaption", "w:tblDescription", "w:tblPrChange")


def _margens_celula(tabela, twips=70):
    tblPr = tabela._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for lado in ("left", "right"):
        m = OxmlElement(f"w:{lado}")
        m.set(qn("w:w"), str(twips))
        m.set(qn("w:type"), "dxa")
        mar.append(m)
    tblPr.insert_element_before(mar, "w:tblLook", "w:tblCaption", "w:tblDescription", "w:tblPrChange")


def _linha_de_cabecalho(row):
    trPr = row._tr.get_or_add_trPr()
    h = OxmlElement("w:tblHeader")
    h.set(qn("w:val"), "true")
    trPr.append(h)


def _borda_inferior(paragrafo):
    pPr = paragrafo._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    b.set(qn("w:val"), "single")
    b.set(qn("w:sz"), "6")
    b.set(qn("w:space"), "1")
    b.set(qn("w:color"), "A6A6A6")
    pBdr.append(b)
    pPr.insert_element_before(pBdr, "w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku",
                              "w:wordWrap", "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE",
                              "w:autoSpaceDN", "w:bidi", "w:adjustRightInd", "w:snapToGrid",
                              "w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents",
                              "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment",
                              "w:textboxTightWrap", "w:outlineLvl", "w:divId", "w:cnfStyle",
                              "w:rPr", "w:sectPr", "w:pPrChange")


def _campo(paragrafo, instrucao, provisorio="1"):
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), instrucao)
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = provisorio
    r.append(t)
    fld.append(r)
    paragrafo._p.append(fld)


def _estilo(doc, nome):
    """O estilo, se o modelo tiver; senão None (quem chama formata à mão)."""
    try:
        return doc.styles[nome]
    except KeyError:
        return None


def _no_corpo(corpo, el):
    """Põe `el` no fim do corpo, antes do `w:sectPr` final. O python-docx procura
    o sectPr varrendo TODOS os filhos do corpo a cada inserção."""
    try:
        ultimo = corpo[-1]          # O(1) no lxml; len(corpo) percorre os filhos
    except IndexError:
        ultimo = None
    if ultimo is not None and ultimo.tag == qn("w:sectPr"):
        ultimo.addprevious(el)
    else:
        corpo.append(el)


def paragrafo(cont, texto="", estilo=None):
    """Parágrafo novo no documento ou numa célula."""
    if hasattr(cont, "_body"):
        from docx.text.paragraph import Paragraph
        corpo = cont.element.body
        el = corpo._new_p()
        _no_corpo(corpo, el)
        p = Paragraph(el, cont._body)
    else:
        p = cont.add_paragraph()
    if texto:
        p.add_run(texto)
    if estilo is not None:
        p.style = estilo
    return p


def nova_tabela(cont, linhas, colunas, largura_cm):
    """Tabela nova no documento ou numa célula. No documento, sem o
    `Document.add_table`, que recalcula a largura da página varrendo o documento
    inteiro a cada tabela."""
    if hasattr(cont, "_body"):
        from docx.oxml.table import CT_Tbl
        from docx.table import Table
        el = CT_Tbl.new_tbl(linhas, colunas, Cm(largura_cm))
        _no_corpo(cont.element.body, el)
        return Table(el, cont._body)
    return cont.add_table(rows=linhas, cols=colunas)


def _tem_texto(p):
    return p._p.find(qn("w:r")) is not None or p._p.find(qn("w:hyperlink")) is not None


def grade(tabela):
    """Células em lista (linha a linha), calculadas UMA vez e só para tabela
    recém-criada (sem mescla: cada linha tem exatamente uma célula por coluna).
    `Table.cell(r, c)` refaz a grade inteira a cada chamada."""
    from docx.table import _Cell
    return [_Cell(tc, tabela) for tr in tabela._tbl.tr_lst for tc in tr.tc_lst]


def _acelerar(doc):
    """Pontos do python-docx que são quadráticos no tamanho do documento: num
    caderno inteiro (600 imagens, 600 tabelas, 27 mil parágrafos) levavam a
    geração para mais de um minuto. Tudo vale só para ESTA instância:

      * imagem repetida é achada comparando o SHA-1 de cada imagem já posta,
        recalculado a cada comparação → aqui, um dicionário;
      * nome de cada imagem (`word/media/imageN`): testa N contra uma lista;
      * id de cada imagem: o próximo é achado varrendo todos os ids do documento.

    (Parágrafo e tabela no corpo: ver `paragrafo` e `nova_tabela`.)
    """
    from docx.image.image import Image as ImagemDocx
    from docx.opc.packuri import PackURI
    imagens = doc.part.package.image_parts
    proximo = [len(imagens)]
    por_sha1 = {}

    def nome_imagem(ext):
        proximo[0] += 1
        return PackURI(f"/word/media/image{proximo[0]}.{ext}")

    def obter_ou_criar(descritor):
        img = ImagemDocx.from_file(descritor)
        parte_img = por_sha1.get(img.sha1)
        if parte_img is None:
            parte_img = por_sha1[img.sha1] = imagens._add_image_part(img)
        return parte_img

    imagens._next_image_partname = nome_imagem
    imagens.get_or_add_image_part = obter_ou_criar

    parte = doc.part
    base = type(parte)

    class _ParteComContador(base):
        @property
        def next_id(self):
            self._doctrack_prox_id += 1
            return self._doctrack_prox_id

    parte._doctrack_prox_id = base.next_id.fget(parte)
    parte.__class__ = _ParteComContador


def _limpar_celula(cell):
    """Uma célula nasce com um parágrafo vazio. Se algo foi escrito depois dele,
    ele sai; e o Word exige que a célula termine num parágrafo (não numa tabela)."""
    tc = cell._tc
    ps = tc.findall(qn("w:p"))
    blocos = [e for e in tc if e.tag in (qn("w:p"), qn("w:tbl"))]
    if len(blocos) > 1 and ps and blocos[0] is ps[0] and not "".join(ps[0].itertext()).strip() \
            and not ps[0].findall(".//" + qn("w:drawing")):
        tc.remove(ps[0])
    if tc[-1].tag == qn("w:tbl"):
        tc.append(OxmlElement("w:p"))


# ══════════════════════════════════════════════════════════════════════════════
# IMAGENS
# ══════════════════════════════════════════════════════════════════════════════

def preparar_imagem(dados, ext, largura_cm):
    """(bytes para o Word, largura_px, altura_px). Reamostra para ~DPI_MAX na
    largura impressa e converte o que o Word não abre (webp) para PNG."""
    from PIL import Image
    im = Image.open(io.BytesIO(dados))
    im.load()
    px_w, px_h = im.size
    ext = (ext or "").lower()
    alvo_px = max(1, int(largura_cm / 2.54 * DPI_MAX))
    converter = ext in (".webp", ".gif")
    if px_w <= alvo_px * 1.15 and not converter:
        return dados, px_w, px_h
    jpeg = ext in (".jpg", ".jpeg")
    if im.mode in ("P", "1", "LA") or (im.mode == "RGBA" and jpeg) or im.mode == "CMYK":
        im = im.convert("RGB" if jpeg or im.mode == "CMYK" else "RGBA")
    if px_w > alvo_px * 1.15:
        im = im.resize((alvo_px, max(1, round(px_h * alvo_px / px_w))), Image.LANCZOS)
    buf = io.BytesIO()
    if jpeg:
        im.save(buf, "JPEG", quality=88)
    else:
        im.save(buf, "PNG")
    return buf.getvalue(), px_w, px_h


# ══════════════════════════════════════════════════════════════════════════════
# HTML → WORD
# ══════════════════════════════════════════════════════════════════════════════

@dataclass
class Figura:
    numero: int
    uuid: str
    pagina: int            # índice da página na exportação (1..n)
    legenda: str


@dataclass
class _Ctx:
    largura_cm: float                  # largura disponível (página ou célula)
    tamanho: object = None             # Pt padrão do texto (tabela usa menor)
    lista: list = field(default_factory=list)   # pilha de [tipo, contador, estilo]
    alinhamento: object = None


class _Escritor:
    """Escreve a árvore de UMA página num contêiner do python-docx (o documento
    ou uma célula — os dois têm add_paragraph/add_table)."""

    def __init__(self, rel, pagina_idx):
        self.rel = rel
        self.doc = rel.doc
        self.pagina_idx = pagina_idx

    # ── blocos ──
    def blocos(self, cont, nos, ctx):
        pendente = []           # inline acumulado até o próximo bloco

        def descarregar():
            if pendente and any(not (isinstance(n, str) and not n.strip()) for n in pendente):
                p = self._paragrafo(cont, ctx)
                self.inline(p, pendente, {}, ctx)
            pendente.clear()

        for n in nos:
            if isinstance(n, str) or n.tag not in _BLOCOS:
                if isinstance(n, No) and n.tag == "img" and self._imagem_sozinha(nos):
                    descarregar()
                    self.figura(cont, n, None, ctx)
                    continue
                pendente.append(n)
                continue
            descarregar()
            self.bloco(cont, n, ctx)
        descarregar()

    def _imagem_sozinha(self, nos):
        reais = [n for n in nos if not (isinstance(n, str) and not n.strip())]
        return len(reais) == 1

    def _paragrafo(self, cont, ctx, estilo=None):
        p = paragrafo(cont, estilo=estilo)
        if ctx.alinhamento is not None:
            p.alignment = ctx.alinhamento
        if ctx.tamanho is not None:          # dentro de tabela: sem respiro extra
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.space_before = Pt(0)
        return p

    def _alinhar(self, p, no):
        est = no.estilo()
        al = {"center": WD_ALIGN_PARAGRAPH.CENTER, "right": WD_ALIGN_PARAGRAPH.RIGHT,
              "justify": WD_ALIGN_PARAGRAPH.JUSTIFY, "left": WD_ALIGN_PARAGRAPH.LEFT}.get(
            (est.get("text-align") or "").lower())
        if al is not None:
            p.alignment = al
        recuo = _cm_de_css(est.get("margin-left") or est.get("padding-left") or "", 0)
        if recuo:
            p.paragraph_format.left_indent = Cm(min(recuo, 10))

    def bloco(self, cont, n, ctx):
        t = n.tag
        if t == "p":
            p = self._paragrafo(cont, ctx)
            self._alinhar(p, n)
            self.inline(p, n.filhos, {}, ctx)
        elif t in ("h1", "h2", "h3", "h4"):
            nivel = {"h1": 2, "h2": 2, "h3": 3, "h4": 4}[t]
            est = _estilo(self.doc, f"Heading {nivel}")
            p = self._paragrafo(cont, ctx, est)
            self._alinhar(p, n)
            self.inline(p, n.filhos, {} if est is not None else {"negrito": True}, ctx)
        elif t in ("ul", "ol"):
            self.lista(cont, n, ctx)
        elif t == "li":                       # li solto (sem ul/ol): vira item com marcador
            self.lista(cont, No("ul", {}), ctx, itens=[n])
        elif t == "table":
            self.tabela(cont, n, ctx)
        elif t == "figure":
            self.figure(cont, n, ctx)
        elif t == "figcaption":
            self._legenda(cont, n.texto().strip(), ctx)
        elif t == "blockquote":
            est = _estilo(self.doc, "Quote")
            for filho in n.filhos:
                if isinstance(filho, No) and filho.tag == "p":
                    p = self._paragrafo(cont, ctx, est)
                    self.inline(p, filho.filhos, {} if est is not None else {"italico": True}, ctx)
                elif isinstance(filho, No):
                    self.bloco(cont, filho, ctx)
        elif t == "pre":
            for linha in n.texto().split("\n"):
                p = self._paragrafo(cont, ctx)
                p.paragraph_format.space_after = Pt(0)
                r = p.add_run(linha)
                r.font.name = "Consolas"
                r.font.size = Pt(9)
        elif t == "hr":
            _borda_inferior(self._paragrafo(cont, ctx))
        else:                                 # div, thead/tbody fora de tabela...
            self.blocos(cont, n.filhos, ctx)

    # ── listas (numeração escrita no texto: independe do modelo do Word) ──
    def lista(self, cont, n, ctx, itens=None):
        tipo = n.tag
        est = n.estilo().get("list-style-type", "")
        try:
            contador = int(n.attrs.get("start") or 1)
        except ValueError:
            contador = 1
        tarefas = "todo-list" in n.classes()
        ctx.lista.append([tipo, contador, est, tarefas])
        for li in itens or [f for f in n.filhos if isinstance(f, No) and f.tag == "li"]:
            self.item(cont, li, ctx)
            ctx.lista[-1][1] += 1
        ctx.lista.pop()

    def _marcador(self, ctx):
        tipo, n, est, tarefas = ctx.lista[-1]
        if tarefas:
            return ""
        prof = len(ctx.lista) - 1
        if tipo == "ul":
            return ["•", "◦", "▪"][prof % 3]
        if est == "lower-alpha" or est == "lower-latin":
            return _alfa(n) + "."
        if est == "upper-alpha" or est == "upper-latin":
            return _alfa(n).upper() + "."
        if est == "lower-roman":
            return _romano(n).lower() + "."
        if est == "upper-roman":
            return _romano(n) + "."
        return f"{n}."

    def item(self, cont, li, ctx):
        prof = len(ctx.lista)
        recuo = Cm(0.63 * prof)
        marcador = self._marcador(ctx)
        primeiro = True
        pendente = []

        def novo_paragrafo():
            p = self._paragrafo(cont, ctx)
            p.paragraph_format.left_indent = recuo
            p.paragraph_format.space_after = Pt(1)
            return p

        def descarregar():
            nonlocal primeiro
            if not pendente:
                return
            p = novo_paragrafo()
            if primeiro and marcador:
                p.paragraph_format.first_line_indent = Cm(-0.5)
                p.add_run(marcador + "\t")
            primeiro = False
            self.inline(p, pendente, {}, ctx)
            pendente.clear()

        for f in li.filhos:
            if isinstance(f, No) and f.tag in ("ul", "ol"):
                descarregar()
                self.lista(cont, f, ctx)
            elif isinstance(f, No) and f.tag == "p":
                descarregar()
                pendente.extend(f.filhos)
                descarregar()
            elif isinstance(f, No) and f.tag in _BLOCOS:
                descarregar()
                self.bloco(cont, f, ctx)
                primeiro = False
            else:
                pendente.append(f)
        descarregar()
        if primeiro:                     # item vazio: mantém o número
            p = novo_paragrafo()
            p.paragraph_format.first_line_indent = Cm(-0.5)
            p.add_run(marcador)

    # ── tabelas ──
    def tabela(self, cont, n, ctx, legenda=None, largura_css=None):
        linhas = []
        cabecalho = set()
        larguras = []
        for f in n.filhos:
            if not isinstance(f, No):
                continue
            if f.tag == "colgroup":
                larguras = [c.estilo().get("width", "") for c in f.filhos
                            if isinstance(c, No) and c.tag == "col"]
            elif f.tag in ("thead", "tbody", "tfoot"):
                for tr in f.filhos:
                    if isinstance(tr, No) and tr.tag == "tr":
                        if f.tag == "thead":
                            cabecalho.add(len(linhas))
                        linhas.append(tr)
            elif f.tag == "tr":
                linhas.append(f)
        if not linhas:
            return

        # Grade com colspan/rowspan.
        ocupado = {}
        celulas = []          # (linha, coluna, rs, cs, no)
        ncols = 0
        for r, tr in enumerate(linhas):
            c = 0
            for td in [x for x in tr.filhos if isinstance(x, No) and x.tag in ("td", "th")]:
                while (r, c) in ocupado:
                    c += 1
                rs = max(1, min(_inteiro(td.attrs.get("rowspan"), 1), len(linhas) - r))
                cs = max(1, min(_inteiro(td.attrs.get("colspan"), 1), 63))
                for dr in range(rs):
                    for dc in range(cs):
                        ocupado[(r + dr, c + dc)] = True
                celulas.append((r, c, rs, cs, td))
                c += cs
            ncols = max(ncols, c)
        ncols = max(ncols, 1)

        largura_total = ctx.largura_cm
        if largura_css:
            w = _cm_de_css(largura_css, ctx.largura_cm)
            if w:
                largura_total = min(w, ctx.largura_cm)
        cols_cm = []
        if larguras and len(larguras) == ncols:
            vals = [_cm_de_css(x, largura_total) for x in larguras]
            if all(vals):
                soma = sum(vals)
                cols_cm = [v * largura_total / soma for v in vals]
        if not cols_cm:
            cols_cm = [largura_total / ncols] * ncols

        t = nova_tabela(cont, len(linhas), ncols, sum(cols_cm))
        t.alignment = WD_TABLE_ALIGNMENT.LEFT
        t.autofit = False
        _bordas_tabela(t)
        _margens_celula(t)
        for i, w in enumerate(cols_cm):
            t.columns[i].width = Cm(w)
        g = grade(t)
        for k, cell in enumerate(g):
            cell.width = Cm(cols_cm[k % ncols])
        linhas_xml = t.rows
        for r in sorted(cabecalho):
            if r == 0 or (r - 1) in cabecalho:
                _linha_de_cabecalho(linhas_xml[r])

        # Só a célula de origem de cada mescla é escrita; as posições cobertas
        # por outra mescla nunca são origem (a grade acima garante), então a
        # lista calculada antes das mesclas continua valendo para elas.
        for r, c, rs, cs, td in celulas:
            cell = g[r * ncols + c]
            if rs > 1 or cs > 1:
                cell = cell.merge(g[(r + rs - 1) * ncols + (c + cs - 1)])
            est = td.estilo()
            fundo = cor_hex(est.get("background-color") or est.get("background"))
            th = td.tag == "th"
            if fundo:
                _sombrear_celula(cell, fundo)
            elif th:
                _sombrear_celula(cell, CINZA_TH)
            al = {"center": WD_ALIGN_PARAGRAPH.CENTER, "right": WD_ALIGN_PARAGRAPH.RIGHT,
                  "justify": WD_ALIGN_PARAGRAPH.JUSTIFY}.get((est.get("text-align") or "").lower())
            sub = _Ctx(largura_cm=max(1.0, sum(cols_cm[c:c + cs]) - 0.25), tamanho=FONTE_TABELA,
                       alinhamento=al)
            filhos = td.filhos
            if th:
                filhos = [No("strong", {})]
                filhos[0].filhos = td.filhos
            self.blocos(cell, filhos, sub)
            _limpar_celula(cell)
        if legenda:
            self._legenda(cont, legenda, ctx, prefixo="Tabela")
        else:
            # Respiro depois da tabela, senão a próxima tabela "gruda" nesta.
            p = self._paragrafo(cont, ctx)
            p.paragraph_format.space_after = Pt(2)

    # ── figure (imagem ou tabela com legenda) ──
    def figure(self, cont, n, ctx):
        legenda = ""
        for f in n.filhos:
            if isinstance(f, No) and f.tag == "figcaption":
                legenda = re.sub(r"\s+", " ", f.texto()).strip()
        if "table" in n.classes():
            for f in n.filhos:
                if isinstance(f, No) and f.tag == "table":
                    self.tabela(cont, f, ctx, legenda=legenda, largura_css=n.estilo().get("width"))
            return
        img = next((f for f in n.filhos if isinstance(f, No) and f.tag == "img"), None)
        if img is None:
            self.blocos(cont, [f for f in n.filhos if not (isinstance(f, No) and f.tag == "figcaption")], ctx)
            return
        self.figura(cont, img, n, ctx, legenda=legenda)

    def _largura_imagem(self, img, figura, ctx, px_w):
        larg = None
        for fonte in (img.estilo().get("width"), figura.estilo().get("width") if figura else None,
                      img.attrs.get("width")):
            if fonte and fonte != "auto":
                larg = _cm_de_css(fonte, ctx.largura_cm)
                if larg:
                    break
        if not larg:
            larg = px_w * 2.54 / 96
        return max(0.5, min(larg, ctx.largura_cm))

    def figura(self, cont, img, figura, ctx, legenda=""):
        uuid = img.attrs.get("data-arquivo")
        dados = self.rel.imagem_de(uuid) if uuid else None
        p = self._paragrafo(cont, ctx)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        if not dados:
            p.add_run("[imagem indisponível]").italic = True
            return
        try:
            from PIL import Image
            with Image.open(io.BytesIO(dados[0])) as im:
                px_w = im.size[0]
            larg = self._largura_imagem(img, figura, ctx, px_w)
            conteudo, _, _ = preparar_imagem(dados[0], dados[1], larg)
            p.add_run().add_picture(io.BytesIO(conteudo), width=Cm(larg))
        except Exception:                       # imagem corrompida não derruba o relatório
            p.add_run("[imagem que não pôde ser lida]").italic = True
            return
        self.rel.n_figura += 1
        # Só a legenda escrita (figcaption). O `alt` não serve: no conteúdo vindo
        # do OneNote ele é o OCR automático da imagem.
        self.rel.figuras.append(Figura(self.rel.n_figura, uuid, self.pagina_idx, legenda))
        self._legenda(cont, legenda, ctx, prefixo=f"Figura {self.rel.n_figura}")

    def imagem_em_linha(self, p, img, ctx):
        uuid = img.attrs.get("data-arquivo")
        dados = self.rel.imagem_de(uuid) if uuid else None
        if not dados:
            p.add_run("[imagem]").italic = True
            return
        try:
            from PIL import Image
            with Image.open(io.BytesIO(dados[0])) as im:
                px_w = im.size[0]
            larg = self._largura_imagem(img, None, ctx, px_w)
            conteudo, _, _ = preparar_imagem(dados[0], dados[1], larg)
            p.add_run().add_picture(io.BytesIO(conteudo), width=Cm(larg))
            self.rel.sem_legenda.append((uuid, self.pagina_idx))
        except Exception:
            p.add_run("[imagem]").italic = True

    def _legenda(self, cont, texto, ctx, prefixo=None):
        if not texto and not prefixo:
            return
        est = _estilo(self.doc, "Caption")
        p = self._paragrafo(cont, ctx, est)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(f"{prefixo} — {texto}" if prefixo and texto else (prefixo or texto))
        if est is None:
            r.italic = True
            r.font.size = Pt(9)

    # ── inline ──
    def inline(self, p, nos, fmt, ctx):
        for n in nos:
            if isinstance(n, str):
                texto = re.sub(r"\s+", " ", n)
                if not texto:
                    continue
                if not _tem_texto(p):
                    if not texto.strip():
                        continue
                    texto = texto.lstrip()
                self._run(p, texto, fmt, ctx)
                continue
            t = n.tag
            novo = dict(fmt)
            if t in ("strong", "b"):
                novo["negrito"] = True
            elif t in ("i", "em"):
                novo["italico"] = True
            elif t == "u":
                novo["sublinhado"] = True
            elif t in ("s", "del"):
                novo["tachado"] = True
            elif t == "sub":
                novo["subscrito"] = True
            elif t == "sup":
                novo["sobrescrito"] = True
            elif t == "code":
                novo["mono"] = True
            elif t == "br":
                if _tem_texto(p):
                    p.add_run().add_break()
                else:
                    p.add_run().add_break()
                continue
            elif t == "img":
                self.imagem_em_linha(p, n, ctx)
                continue
            elif t == "input":
                self._run(p, "☑ " if "checked" in n.attrs else "☐ ", fmt, ctx)
                continue
            elif t == "a":
                arq = n.attrs.get("data-arquivo")
                href = n.attrs.get("href") or ""
                if arq:
                    destino = self.rel.link_de(arq)
                    if destino:
                        novo["link"] = quote(destino, safe="/")
                    else:
                        novo["nao_incluido"] = True
                elif href.lower().startswith(("http://", "https://", "mailto:")):
                    novo["link"] = href
            for c in n.classes():
                if c in REALCES:
                    novo["realce"] = REALCES[c]
                elif c in CANETAS:
                    novo["cor"] = CANETAS[c]
                elif c in TAMANHOS:
                    novo["escala"] = TAMANHOS[c]
            est = n.estilo()
            if est.get("color"):
                cor = cor_hex(est["color"])
                if cor and cor != "000000":
                    novo["cor"] = cor
            if est.get("background-color"):
                fundo = cor_hex(est["background-color"])
                if fundo and fundo != "FFFFFF":
                    novo["fundo"] = fundo
            if est.get("font-size"):
                m = re.fullmatch(r"([\d.]+)(px|pt)", est["font-size"].strip().lower())
                if m:
                    v = float(m.group(1))
                    novo["pt"] = v * 0.75 if m.group(2) == "px" else v
            if t == "mark" and "realce" not in novo:
                novo["realce"] = WD_COLOR_INDEX.YELLOW
            self.inline(p, n.filhos, novo, ctx)
            if t == "a" and novo.get("nao_incluido"):
                r = p.add_run(" (arquivo fora do pacote)")
                r.italic = True
                r.font.size = Pt(8)
                r.font.color.rgb = RGBColor(0x80, 0x80, 0x80)

    def _run(self, p, texto, fmt, ctx):
        r = p.add_run(texto)
        f = r.font
        if fmt.get("negrito"):
            r.bold = True
        if fmt.get("italico"):
            r.italic = True
        if fmt.get("sublinhado"):
            r.underline = True
        if fmt.get("tachado"):
            f.strike = True
        if fmt.get("subscrito"):
            f.subscript = True
        if fmt.get("sobrescrito"):
            f.superscript = True
        if fmt.get("mono"):
            f.name = "Consolas"
        base = ctx.tamanho.pt if ctx.tamanho is not None else None
        if fmt.get("pt"):
            f.size = Pt(min(max(fmt["pt"], 6), 28))
        elif fmt.get("escala"):
            f.size = Pt(round((base or 10.5) * fmt["escala"], 1))
        elif ctx.tamanho is not None:
            f.size = ctx.tamanho
        if fmt.get("cor"):
            f.color.rgb = RGBColor.from_string(fmt["cor"])
        if fmt.get("realce") is not None:
            f.highlight_color = fmt["realce"]
        if fmt.get("fundo"):
            _sombrear_run(r, fmt["fundo"])
        if fmt.get("link"):
            if not fmt.get("cor"):
                f.color.rgb = AZUL_LINK
            r.underline = True
            r_id = p.part.relate_to(fmt["link"], RT.HYPERLINK, is_external=True)
            h = OxmlElement("w:hyperlink")
            h.set(qn("r:id"), r_id)
            r._r.addprevious(h)
            h.append(r._r)
        return r


def _inteiro(v, padrao):
    try:
        return int(v)
    except (TypeError, ValueError):
        return padrao


def _alfa(n):
    s = ""
    while n > 0:
        n, r = divmod(n - 1, 26)
        s = chr(97 + r) + s
    return s or "a"


def _romano(n):
    tabela = [(1000, "M"), (900, "CM"), (500, "D"), (400, "CD"), (100, "C"), (90, "XC"),
              (50, "L"), (40, "XL"), (10, "X"), (9, "IX"), (5, "V"), (4, "IV"), (1, "I")]
    s = ""
    for v, r in tabela:
        while n >= v:
            s += r
            n -= v
    return s or "I"


# ══════════════════════════════════════════════════════════════════════════════
# CONCLUSÃO (para a tabela de testes)
# ══════════════════════════════════════════════════════════════════════════════

def conclusao_de(html_texto, limite=320):
    """Texto sob o título que contém "CONCLUSÃO" (h2–h4), até o próximo título
    do mesmo nível ou acima. É o "Conclusão" do modelo de registro do PDR."""
    raiz = arvore(html_texto)
    nivel_achado = None
    partes = []
    for n in raiz.filhos:
        if isinstance(n, No) and n.tag in ("h1", "h2", "h3", "h4"):
            nivel = int(n.tag[1])
            if nivel_achado is not None and nivel <= nivel_achado:
                break
            if nivel_achado is None and "conclus" in _sem_acento(n.texto()):
                nivel_achado = nivel
                continue
        if nivel_achado is not None:
            partes.append(n if isinstance(n, str) else n.texto())
    texto = re.sub(r"\s+", " ", " ".join(partes)).strip()
    return texto if len(texto) <= limite else texto[:limite - 1].rstrip() + "…"


# ══════════════════════════════════════════════════════════════════════════════
# O RELATÓRIO
# ══════════════════════════════════════════════════════════════════════════════

@dataclass
class PaginaRelatorio:
    titulo: str
    secao: str
    tipo: str
    html: str
    versao: int
    data_pagina: datetime | None = None
    ficha: dict = field(default_factory=dict)
    resp_teste: str = ""
    resp_execucao: str = ""
    autores: str = ""


@dataclass
class ArquivoRelatorio:
    caminho: str
    tamanho: int
    origem: str


def _data(valor):
    if not valor:
        return "—"
    if isinstance(valor, str):
        try:
            return datetime.strptime(valor[:10], "%Y-%m-%d").strftime("%d/%m/%Y")
        except ValueError:
            return valor
    return valor.strftime("%d/%m/%Y")


def tamanho_legivel(b):
    b = int(b or 0)
    if b < 1024:
        return f"{b} B"
    if b < 1048576:
        return f"{b / 1024:.0f} KB"
    if b < 1073741824:
        return f"{b / 1048576:.1f} MB".replace(".", ",")
    return f"{b / 1073741824:.2f} GB".replace(".", ",")


class Relatorio:
    """Uso: r = Relatorio(...); dados = r.gerar(paginas, arquivos) → bytes do .docx.
    Depois de gerar, `r.figuras` tem a numeração das figuras (para `imagens/`)."""

    def __init__(self, *, caderno, escopo, emitido_por, emitido_em=None,
                 imagem_de=lambda uuid: None, link_de=lambda uuid: None):
        self.caderno = caderno
        self.escopo = escopo
        self.emitido_por = emitido_por
        self.emitido_em = emitido_em or datetime.now()
        self.imagem_de = imagem_de        # uuid → (bytes, ext) | None
        self.link_de = link_de            # uuid → caminho no zip | None
        self.figuras = []
        self.sem_legenda = []
        self.n_figura = 0
        self.doc = None

    # ── documento base ──
    def _novo_documento(self):
        if os.path.isfile(MODELO):
            doc = Document(MODELO)
            corpo = doc.element.body
            for el in list(corpo):
                if el.tag != qn("w:sectPr"):
                    corpo.remove(el)
            return doc, True
        doc = Document()
        sec = doc.sections[0]
        sec.orientation = WD_ORIENT.PORTRAIT
        sec.page_width, sec.page_height = Cm(21), Cm(29.7)
        sec.left_margin = sec.right_margin = Cm(2)
        sec.top_margin, sec.bottom_margin = Cm(2), Cm(2)
        normal = doc.styles["Normal"]
        normal.font.name = "Calibri"
        normal.font.size = Pt(10.5)
        normal.paragraph_format.space_after = Pt(4)
        for nivel, tam in ((1, 16), (2, 13), (3, 11.5), (4, 10.5)):
            est = doc.styles[f"Heading {nivel}"]
            est.font.size = Pt(tam)
            est.paragraph_format.space_before = Pt(14 if nivel <= 2 else 10)
            est.paragraph_format.space_after = Pt(4)
        return doc, False

    def _cabecalho_rodape(self, doc):
        sec = doc.sections[0]
        cab = sec.header.paragraphs[0]
        cab.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = cab.add_run(f"{self.caderno} · {self.escopo}")
        r.font.size = Pt(8)
        r.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)
        rod = sec.footer.paragraphs[0]
        rod.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = rod.add_run(f"Gerado pelo DocTrack em {self.emitido_em:%d/%m/%Y %H:%M} · página ")
        r.font.size = Pt(8)
        _campo(rod, "PAGE")
        r = rod.add_run(" de ")
        r.font.size = Pt(8)
        _campo(rod, "NUMPAGES")

    def _largura_util(self, doc):
        s = doc.sections[-1]
        return (s.page_width - s.left_margin - s.right_margin) / 360000      # EMU → cm

    def _tabela_simples(self, linhas, larguras, cabecalho=True, tamanho=FONTE_TABELA, vazio="—"):
        doc = self.doc
        t = nova_tabela(doc, len(linhas), len(larguras), sum(larguras))
        t.autofit = False
        _bordas_tabela(t)
        _margens_celula(t)
        for i, w in enumerate(larguras):
            t.columns[i].width = Cm(w)
        g = grade(t)
        for r, linha in enumerate(linhas):
            for c, valor in enumerate(linha):
                cell = g[r * len(larguras) + c]
                cell.width = Cm(larguras[c])
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                run = p.add_run(str(valor) if valor not in (None, "") else vazio)
                run.font.size = tamanho
                if cabecalho and r == 0:
                    run.bold = True
                    _sombrear_celula(cell, CINZA_TH)
        if cabecalho:
            _linha_de_cabecalho(t.rows[0])
        return t

    def _ficha(self, pg, largura):
        f = pg.ficha or {}
        linhas = [
            ("Projeto", f.get("projeto") or "—", "Número do teste", f.get("numero") or "—"),
            ("Responsável pelo teste", pg.resp_teste or "—", "Responsável pela execução", pg.resp_execucao or "—"),
            ("Início", _data(f.get("inicio")), "Término", _data(f.get("termino"))),
            ("Resultado", RESULTADO_ROTULO.get(f.get("resultado") or "", "—"), "", ""),
        ]
        cols = [largura * 0.2, largura * 0.3, largura * 0.2, largura * 0.3]
        t = nova_tabela(self.doc, len(linhas), 4, sum(cols))
        t.autofit = False
        _bordas_tabela(t, "D1D5DB")
        _margens_celula(t)
        for i, w in enumerate(cols):
            t.columns[i].width = Cm(w)
        g = grade(t)
        for r, linha in enumerate(linhas):
            for c, valor in enumerate(linha):
                cell = g[r * 4 + c]
                cell.width = Cm(cols[c])
                run = cell.paragraphs[0].add_run(valor)
                cell.paragraphs[0].paragraph_format.space_after = Pt(0)
                run.font.size = FONTE_TABELA
                if c % 2 == 0:
                    run.bold = True
                    _sombrear_celula(cell, "F9FAFB")
        paragrafo(self.doc).paragraph_format.space_after = Pt(2)

    def gerar(self, paginas, arquivos, progresso=None):
        doc, do_modelo = self._novo_documento()
        _acelerar(doc)
        self.doc = doc
        if not do_modelo:
            self._cabecalho_rodape(doc)
        largura = self._largura_util(doc)
        registros = [p for p in paginas if p.tipo == "registro_teste"]

        # ── capa ──
        est_titulo = _estilo(doc, "Title")
        p = paragrafo(doc, estilo=est_titulo)
        r = p.add_run("Relatório de testes" if registros else "Relatório do caderno")
        if est_titulo is None:
            r.bold, r.font.size = True, Pt(24)
        p = paragrafo(doc)
        r = p.add_run(self.caderno)
        r.bold, r.font.size = True, Pt(15)
        inicios = [p.ficha.get("inicio") for p in registros if p.ficha.get("inicio")]
        terminos = [p.ficha.get("termino") for p in registros if p.ficha.get("termino")]
        if inicios or terminos:
            periodo = f"{_data(min(inicios) if inicios else None)} a {_data(max(terminos) if terminos else None)}"
        else:
            datas = [p.data_pagina for p in paginas if p.data_pagina]
            periodo = f"{_data(min(datas))} a {_data(max(datas))}" if datas else "—"
        self._tabela_simples([
            ("Escopo", self.escopo),
            ("Período", periodo),
            ("Páginas", f"{len(paginas)} ({len(registros)} registros de teste)"),
            ("Arquivos no pacote", f"{len(arquivos)} ({tamanho_legivel(sum(a.tamanho for a in arquivos))})"),
            ("Emitido por", self.emitido_por),
            ("Emitido em", f"{self.emitido_em:%d/%m/%Y %H:%M}"),
            ("Origem", "DocTrack · Cadernos de laboratório"),
        ], [largura * 0.3, largura * 0.7], cabecalho=False, tamanho=Pt(10))
        paragrafo(doc)

        # ── testes efetuados ──
        paragrafo(doc, "Testes efetuados", _estilo(doc, "Heading 1"))
        if registros:
            linhas = [("Nº", "Título", "Início – término", "Responsáveis", "Resultado", "Conclusão")]
            for pg in registros:
                f = pg.ficha or {}
                resp = " / ".join(x for x in (pg.resp_teste, pg.resp_execucao) if x) or "—"
                datas = (f"{_data(f.get('inicio'))} – {_data(f.get('termino'))}"
                         if f.get("inicio") or f.get("termino") else "—")
                linhas.append((f.get("numero") or "—", pg.titulo, datas, resp,
                               RESULTADO_ROTULO.get(f.get("resultado") or "", "—"),
                               conclusao_de(pg.html) or "—"))
            self._tabela_simples(linhas, [largura * x for x in (0.07, 0.25, 0.15, 0.17, 0.11, 0.25)],
                                 tamanho=Pt(8.5))
        else:
            paragrafo(doc, "Nenhum registro de teste nas páginas exportadas.")

        # ── páginas ──
        for i, pg in enumerate(paginas, 1):
            if progresso:
                progresso(i - 1, len(paginas), f"Página {i} de {len(paginas)}: {pg.titulo}")
            paragrafo(doc).add_run().add_break(WD_BREAK.PAGE)
            paragrafo(doc, pg.titulo or "Sem título", _estilo(doc, "Heading 1"))
            meta = [pg.secao, _data(pg.data_pagina)]
            if pg.autores:
                meta.append(pg.autores)
            meta.append(f"versão {pg.versao}")
            p = paragrafo(doc)
            r = p.add_run(" · ".join(x for x in meta if x and x != "—"))
            r.font.size, r.font.color.rgb = Pt(8.5), RGBColor(0x6B, 0x72, 0x80)
            if pg.tipo == "registro_teste":
                self._ficha(pg, largura)
            raiz = arvore(pg.html)
            if not any(isinstance(n, No) or (isinstance(n, str) and n.strip()) for n in raiz.filhos):
                paragrafo(doc, "(página sem conteúdo)").runs[0].italic = True
                continue
            _Escritor(self, i).blocos(doc, raiz.filhos, _Ctx(largura_cm=largura))

        # ── arquivos ──
        paragrafo(doc).add_run().add_break(WD_BREAK.PAGE)
        paragrafo(doc, "Arquivos do pacote", _estilo(doc, "Heading 1"))
        if arquivos:
            linhas = [("Arquivo", "Tamanho", "Origem")]
            linhas += [(a.caminho, tamanho_legivel(a.tamanho), a.origem) for a in arquivos]
            self._tabela_simples(linhas, [largura * 0.62, largura * 0.13, largura * 0.25], tamanho=Pt(8.5))
        else:
            paragrafo(doc, "Nenhum arquivo além do relatório e das imagens.")
        paragrafo(doc, "A integridade de cada arquivo pode ser conferida pelo SHA-256 listado no "
                          "LEIA-ME.txt do pacote.").runs[0].font.size = Pt(8.5)

        # ── assinaturas ──
        paragrafo(doc, "Assinaturas", _estilo(doc, "Heading 1"))
        self._tabela_simples([("", "Nome", "Data", "Assinatura"),
                              ("Elaborado por", "", "", ""),
                              ("Revisado por", "", "", "")],
                             [largura * 0.2, largura * 0.35, largura * 0.15, largura * 0.3],
                             tamanho=Pt(9), vazio="")
        for row in self.doc.tables[-1].rows[1:]:
            row.height = Cm(1.1)

        buf = io.BytesIO()
        doc.save(buf)
        return buf.getvalue()
