"""
tempos/routes.py — Blueprint do módulo Tempos (apontamento de horas).

Montado sob /tempos no servidor mestre. Reutiliza autenticação, audit log e
event_bus do mestre.

## Acesso

Tempo é trabalho de todo mundo: **admin, gestor e técnico** apontam. `leitura` é
perfil de auditoria externa (ver o seed em servidor.py) e fica de fora — não
executa entregas e não precisa saber quem gastou quanto em quê.

Duas camadas, e a segunda é nova nesta base:

  1. `@so_membros` — o gate de perfil de sempre. `require_role` já embute o
     `@jwt_required()` (auth.py:57), então NUNCA empilhar os dois.
  2. `_meu_ou_403` / `_query_escopada` — o gate de DONO. Os módulos vizinhos são
     todos gated só por perfil; aqui um técnico não alcança o apontamento de
     outro técnico nem para ler, e gestor/admin enxergam o time inteiro.

O gate de dono é helper e não decorator de propósito: ele precisa da linha já
carregada, e todo decorator de auth.py é stateless (identidade e papel). Fazer um
deles consultar o banco esconderia uma query dentro de um decorator — exatamente
o tipo de comportamento invisível que os comentários deste repositório existem
para evitar.

## Rotas

  GET    /tempos/                              página (folha de ponto + relatórios)
  GET    /tempos/api/meta                      vocabulários, alvos recentes, limites
  GET    /tempos/api/alvos                     busca de alvo por tipo (picker)
  GET    /tempos/api/atual                     cronômetro do usuário + servidor_agora
  POST   /tempos/api/iniciar                   abre intervalo (servidor carimba)
  POST   /tempos/api/atual/pausar              fecha com encerrado_por='pausa'
  POST   /tempos/api/atual/parar               fecha com encerrado_por='usuario'
  POST   /tempos/api/<id>/retomar              novo trecho da mesma sessão (só o dono)
  GET    /tempos/api/apontamentos              lista (de, ate, user_id, ref_tipo,
                                               ref_id, projeto_id, q)
  POST   /tempos/api/apontamentos              lançamento manual/retroativo
  GET    /tempos/api/apontamentos/<id>         detalhe
  PUT    /tempos/api/apontamentos/<id>         edita
  DELETE /tempos/api/apontamentos/<id>         exclui
  GET    /tempos/api/resumo                    totais (por=dia|alvo|projeto|pessoa)
  GET    /tempos/api/por-ref                   total de um alvo — para as fichas dos
                                               outros módulos
  GET    /tempos/api/export/apontamentos.csv   CSV (`;` + BOM)
  GET    /tempos/api/export/tempos.xlsx        XLSX (2 abas)
"""
import csv
import io
from datetime import date, datetime, timedelta

from flask import Blueprint, jsonify, render_template, request, send_file
from flask_jwt_extended import get_jwt_identity
from sqlalchemy.exc import IntegrityError

from models import db, User, Projeto, Entregavel, MissaoCartao, Equipamento
from auth import require_role, log_action, get_client_ip, negar_acesso
from event_bus import EventType

from . import core
from .models import (
    Apontamento, REF_TIPOS_TEMPO, REF_TIPOS_LABELS, LIMITE_CRONOMETRO_H,
    LIMITE_APONTAMENTO_H, RETROATIVO_DIAS, JANELA_RETOMADA_H, fmt_dur,
)

tempos_bp = Blueprint(
    "tempos", __name__,
    url_prefix="/tempos",
    template_folder="templates",
)

# CSS/JS ficam em static/ do mestre (1º nível) de propósito: _static_version() só
# varre esse nível, então o cache-busting funciona sem token próprio. O PDR
# precisou de um ASSET_V só porque seus assets ficam dentro do pacote.

_rt = {"socketio": None, "publish_event": None, "AuditLog": None, "EventType": None}


def init_realtime(socketio, publish_event, AuditLog, EventType):
    _rt.update(socketio=socketio, publish_event=publish_event,
               AuditLog=AuditLog, EventType=EventType)


def current_user():
    return User.query.filter_by(email=get_jwt_identity()).first()


def _emit(event_type, payload):
    """Tempo real best-effort — a gravação já foi feita (molde do custos/_emit).

    Só é chamado nas mutações de FOLHA (manual, edição, exclusão). Iniciar,
    pausar, parar e retomar NÃO emitem: publish_event grava uma linha em
    audit_logs por evento (event_bus.py:220), e 20 pessoas × 6 ciclos/dia × 2
    eventos são ~60 mil linhas por ano de puro ruído de UI. A sincronia entre
    abas do próprio usuário sai por localStorage + evento `storage`, que custa
    zero no servidor e ainda funciona no config.html, que não carrega socket.io.
    """
    if _rt["socketio"] and _rt["publish_event"]:
        u = current_user()
        try:
            _rt["publish_event"](event_type, payload,
                                 user_id=u.id if u else None,
                                 user_email=u.email if u else None,
                                 db=db, AuditLog=_rt["AuditLog"],
                                 socketio=_rt["socketio"])
        except Exception:
            pass


# Tempo é trabalho de todo mundo — técnico para cima. Um só nome para o gate,
# para que afrouxá-lo por engano numa rota isolada seja visível na revisão.
def so_membros(fn):
    return require_role("admin", "gestor", "tecnico")(fn)


def _e_gestao(u):
    return bool(u and u.role in ("admin", "gestor"))


# ── GATE DE DONO ──────────────────────────────────────────────────────────────

def _meu_ou_403(ap):
    """(apontamento, None) ou (None, resposta 403).

    Regra única do módulo: o dono mexe no que é dele; admin e gestor enxergam e
    corrigem o de todo mundo. Técnico não alcança o alheio nem para ler — um
    apontamento diz o que a pessoa estava fazendo e quando, e isso não é dado de
    time.

    403 e não 404: o 404 esconderia a existência da linha, mas id sequencial de
    apontamento não é segredo e a mensagem honesta economiza um chamado. O corpo
    do 403 não devolve NADA da linha.
    """
    u = current_user()
    if u and (ap.user_id == u.id or _e_gestao(u)):
        return ap, None
    return None, negar_acesso(f"apontamento #{ap.id} de outro usuário",
                              "Este apontamento é de outro usuário")


def _query_escopada():
    """Query base compartilhada por lista, resumo e os dois exports.

    Compartilhada pelo mesmo motivo do _composicoes_filtradas de
    custos/routes.py:179 — o filtro tem que ser o mesmo nos quatro lugares. Se o
    export ficar mais frouxo que a lista, o vazamento sai pela planilha.
    """
    u = current_user()
    q = Apontamento.query
    if _e_gestao(u):
        alvo = (request.args.get("user_id") or "").strip()
        if alvo.isdigit():
            q = q.filter(Apontamento.user_id == int(alvo))
    else:
        # IGNORA o user_id do cliente de propósito, em vez de devolver 403: o
        # front é o mesmo para todos os perfis e sempre manda o parâmetro. E o
        # filtro é INCONDICIONAL — escrever "if not alvo: filtra por mim" é
        # exatamente o IDOR que o teste de listagem escopada cobre.
        q = q.filter(Apontamento.user_id == u.id)
    return q


def _aplicar_filtros(q):
    """Janela e recortes comuns. Devolve (query, erro_pronto)."""
    de = core.parse_dt((request.args.get("de") or "") + "T00:00") \
        or core.parse_dt(request.args.get("de"))
    ate = core.parse_dt((request.args.get("ate") or "") + "T23:59:59") \
        or core.parse_dt(request.args.get("ate"))
    if de:
        q = q.filter(Apontamento.inicio >= de)
    if ate:
        q = q.filter(Apontamento.inicio <= ate)

    ref_tipo = (request.args.get("ref_tipo") or "").strip()
    if ref_tipo:
        if ref_tipo not in REF_TIPOS_TEMPO:
            return None, (jsonify({"erro": f"alvo inválido: {ref_tipo}"}), 400)
        q = q.filter(Apontamento.ref_tipo == ref_tipo)

    ref_id = (request.args.get("ref_id") or "").strip()
    if ref_id.isdigit():
        q = q.filter(Apontamento.ref_id == int(ref_id))

    projeto_id = (request.args.get("projeto_id") or "").strip()
    if projeto_id.isdigit():
        q = q.filter(Apontamento.projeto_id == int(projeto_id))

    termo = (request.args.get("q") or "").strip()
    if termo:
        like = f"%{termo}%"
        q = q.filter(db.or_(Apontamento.titulo.ilike(like),
                            Apontamento.nota.ilike(like)))
    return q, None


def _lista_filtrada():
    # Cura antes de listar, e para TODO MUNDO — não só para quem está lendo.
    # A cura de /atual é por dono, então um cronômetro esquecido de outra pessoa
    # continuava aberto na folha do time até a tarefa diária passar, exibindo
    # "40h em curso" no meio do relatório do gestor. É a mesma consulta indexada
    # (fim IS NULL AND inicio < corte) e, no caso comum, não escreve nada.
    # Fechar linha alheia aqui não é ação de usuário: é a mesma manutenção que o
    # agendador faria, marcada como 'guarda' e visível para o dono corrigir.
    try:
        core.encerrar_abandonados()
    except Exception:
        db.session.rollback()

    q = _query_escopada()
    q, erro = _aplicar_filtros(q)
    if erro:
        return None, erro
    return q.order_by(Apontamento.inicio.desc()), None


def _campo_dt(d, nome):
    """(valor, erro_pronto) para um horário vindo do corpo de um PUT/POST.

    Distingue "não mandou o campo" de "mandou algo impossível de ler". Sem essa
    distinção o PUT respondia 200 para uma data malformada e mantinha o valor
    antigo: a pessoa via a correção salva e o número seguia errado — o pior par
    possível num módulo cujo produto é o número.

    Valor vazio conta como ausente: é o que o front manda ao limpar um campo.
    """
    if nome not in d:
        return None, None
    bruto = d.get(nome)
    if bruto in (None, ""):
        return None, None
    dt = core.parse_dt(bruto)
    if dt is None:
        return None, (jsonify({"erro": f"Horário de {nome} inválido — "
                                       f"use o formato AAAA-MM-DDTHH:MM"}), 400)
    return dt, None


def _serializar(linhas):
    # core.mapa_refs: 1 query por TIPO de alvo, não 1 por linha. Mora no core
    # porque `core.agregar` também precisa dele — ver a nota lá.
    refs = core.mapa_refs(linhas)
    return [ap.to_dict(ref_info=refs.get(ap.id)) for ap in linhas]


# ── PÁGINA ────────────────────────────────────────────────────────────────────

@tempos_bp.route("/")
@tempos_bp.route("")
def index():
    # A página valida o token no front (igual ao Custos, ao PDR e ao hub); o
    # acesso aos dados é barrado em cada rota /tempos/api/*.
    from servidor import _static_version
    return render_template("tempos/dashboard.html", asset_v=_static_version())


# ── META E ALVOS ──────────────────────────────────────────────────────────────

@tempos_bp.route("/api/meta", methods=["GET"])
@so_membros
def meta():
    u = current_user()
    return jsonify({
        "ref_tipos": REF_TIPOS_TEMPO,
        "ref_tipos_labels": REF_TIPOS_LABELS,
        "agrupamentos": core.AGRUPAMENTOS,
        "recentes": core.alvos_recentes(u.id),
        "eu": {"id": u.id, "nome": u.nome, "role": u.role},
        "gestao": _e_gestao(u),
        # A lista de pessoas só faz sentido para quem pode filtrar por elas.
        "pessoas": ([{"id": x.id, "nome": x.nome}
                     for x in User.query.filter_by(ativo=True)
                     .order_by(User.nome).all()] if _e_gestao(u) else []),
        "limites": {
            "cronometro_h": LIMITE_CRONOMETRO_H,
            "apontamento_h": LIMITE_APONTAMENTO_H,
            "retroativo_dias": RETROATIVO_DIAS,
            "retomada_h": JANELA_RETOMADA_H,
        },
    })


@tempos_bp.route("/api/alvos", methods=["GET"])
@so_membros
def alvos():
    """Busca de alvo para o seletor. Sem isto o front teria que baixar os quatro
    catálogos inteiros só para montar um combo."""
    tipo = (request.args.get("tipo") or "").strip()
    if tipo not in REF_TIPOS_TEMPO or tipo == "livre":
        return jsonify({"erro": "tipo inválido"}), 400
    termo = (request.args.get("q") or "").strip()
    like = f"%{termo}%"
    out = []

    if tipo == "projeto":
        q = Projeto.query.filter_by(ativo=True)
        if termo:
            q = q.filter(Projeto.nome.ilike(like))
        out = [{"id": p.id, "rotulo": p.nome, "extra": p.status or ""}
               for p in q.order_by(Projeto.nome).limit(50).all()]

    elif tipo == "entregavel":
        q = Entregavel.query.join(Projeto).filter(Projeto.ativo.is_(True))
        if termo:
            q = q.filter(Entregavel.tipo.ilike(like))
        out = [{"id": e.id, "rotulo": (e.tipo or "").strip(),
                "extra": e.projeto.nome if e.projeto else ""}
               for e in q.order_by(Entregavel.id.desc()).limit(50).all()]

    elif tipo == "missao_cartao":
        q = MissaoCartao.query.filter(MissaoCartao.concluido.is_(False))
        if termo:
            q = q.filter(MissaoCartao.titulo.ilike(like))
        out = [{"id": c.id, "rotulo": (c.titulo or "").strip(),
                "extra": c.missao.nome if c.missao else ""}
               for c in q.order_by(MissaoCartao.atualizado_em.desc()).limit(50).all()]

    elif tipo == "equipamento":
        q = Equipamento.query.filter_by(ativo=True)
        if termo:
            q = q.filter(db.or_(Equipamento.nome.ilike(like),
                                Equipamento.sku.ilike(like)))
        out = [{"id": e.id, "rotulo": e.nome, "extra": e.sku or ""}
               for e in q.order_by(Equipamento.nome).limit(50).all()]

    return jsonify({"alvos": out})


# ── CRONÔMETRO ────────────────────────────────────────────────────────────────
# As rotas de cronômetro NÃO levam id: resolvem o alvo por (user_id, fim IS NULL),
# o que as torna estruturalmente incapazes de tocar o timer de outra pessoa.

@tempos_bp.route("/api/atual", methods=["GET"])
@so_membros
def atual():
    u = current_user()

    # Cura na leitura: o agendador só passa depois da meia-noite (scheduler.py),
    # então o cronômetro esquecido na sexta só seria cortado no sábado. Aqui a
    # pessoa já vê o número certo na segunda de manhã.
    #
    # Sim, é um GET que escreve. Só quando existe linha aberta passada do corte
    # (caso comum: zero escrita, e esta rota já ia carregar essa linha), e dentro
    # do próprio try: ela roda em TODO carregamento de página dos módulos, e uma
    # cura automática jamais pode derrubar a leitura.
    try:
        core.encerrar_abandonados(user_id=u.id)
    except Exception:
        db.session.rollback()

    corrente = core.em_curso(u.id)
    pausa = core.retomavel(u.id) if not corrente else None
    return jsonify({
        "em_curso": corrente.to_dict() if corrente else None,
        "retomavel": pausa.to_dict() if pausa else None,
        # Âncora do relógio do front. Sem isto o dado estaria certo e o usuário
        # veria "-3h 12m" num navegador adiantado — e abriria um chamado.
        "servidor_agora": datetime.now().isoformat(),
    })


@tempos_bp.route("/api/iniciar", methods=["POST"])
@so_membros
def iniciar():
    u = current_user()
    d = request.get_json(silent=True) or {}

    ref_tipo = (d.get("ref_tipo") or "livre").strip()
    ref_id = d.get("ref_id")
    titulo = d.get("titulo") or ""
    erro = core.validar_ref(ref_tipo, ref_id, titulo)
    if erro:
        return jsonify({"erro": erro}), 400

    # `inicio` do corpo é IGNORADO, não recusado: nenhum remetente legítimo
    # existe, e recusar só produziria um 400 confuso. O servidor é o relógio.
    #
    # `core.abrir` fica DENTRO do try: quem levanta o IntegrityError é o
    # `flush()` que ele faz para resolver o sessao_id, não o commit daqui. Com o
    # try só em volta do commit, a corrida escapava como 500.
    try:
        ap = core.abrir(u, ref_tipo=ref_tipo, ref_id=ref_id, titulo=titulo,
                        nota=d.get("nota") or "", projeto_id=d.get("projeto_id"),
                        origem="cronometro")
        db.session.commit()
    except IntegrityError:
        # ÚNICO ponto do módulo onde IntegrityError é esperado e não bug: com 8
        # threads num processo (render.yaml), clique duplo em "Iniciar" é corrida
        # real e quem resolve é o índice parcial uq_tempo_em_curso. Não alargar
        # este except "para ficar seguro" — ele esconderia erro de verdade.
        db.session.rollback()
        corrente = core.em_curso(u.id)
        return jsonify({"erro": "Você já tem um cronômetro em curso",
                        "em_curso": corrente.to_dict() if corrente else None}), 409

    return jsonify(ap.to_dict()), 201


def _fechar_corrente(motivo):
    u = current_user()
    ap = core.em_curso(u.id)
    if not ap:
        return jsonify({"erro": "Nenhum cronômetro em curso"}), 404
    core.fechar(ap, datetime.now(), motivo)
    db.session.commit()
    return jsonify(ap.to_dict())


@tempos_bp.route("/api/atual/pausar", methods=["POST"])
@so_membros
def pausar():
    return _fechar_corrente("pausa")


@tempos_bp.route("/api/atual/parar", methods=["POST"])
@so_membros
def parar():
    return _fechar_corrente("usuario")


@tempos_bp.route("/api/<int:ap_id>/retomar", methods=["POST"])
@so_membros
def retomar(ap_id):
    """Só o dono, SEM exceção para gestão.

    Diferente de ler e corrigir, retomar *cria* linha nova. Um gestor retomando o
    apontamento alheio ou personificaria a pessoa ou reatribuiria o trabalho em
    silêncio. Supervisão é ler e corrigir; não é autorar.
    """
    u = current_user()
    ap = Apontamento.query.get_or_404(ap_id)
    if ap.user_id != u.id:
        return negar_acesso(f"retomar o apontamento #{ap.id} de outro usuário",
                            "Só dá para retomar o próprio apontamento")
    if ap.fim is None:
        return jsonify({"erro": "Este apontamento ainda está em curso"}), 400

    try:
        # Dentro do try pelo mesmo motivo de /iniciar: o flush de core.abrir é
        # que colide com o índice parcial, não o commit.
        novo = core.retomar(ap, u)
        db.session.commit()
    except IntegrityError:
        db.session.rollback()
        corrente = core.em_curso(u.id)
        return jsonify({"erro": "Você já tem um cronômetro em curso",
                        "em_curso": corrente.to_dict() if corrente else None}), 409
    return jsonify(novo.to_dict()), 201


# ── FOLHA (CRUD) ──────────────────────────────────────────────────────────────

@tempos_bp.route("/api/apontamentos", methods=["GET"])
@so_membros
def listar():
    q, erro = _lista_filtrada()
    if erro:
        return erro
    linhas = q.limit(1000).all()
    # core.segundos, e não duracao_s: a linha aberta vale o relógio correndo, que
    # é o que a tabela exibe. Somar duracao_s aqui faria o total no topo não bater
    # com as linhas logo abaixo.
    total = sum(core.segundos(a) for a in linhas)
    return jsonify({
        "apontamentos": _serializar(linhas),
        "total_s": total,
        "total": fmt_dur(total),
    })


@tempos_bp.route("/api/apontamentos", methods=["POST"])
@so_membros
def criar():
    """Lançamento manual/retroativo — o único caminho, junto do PUT, em que o
    cliente afirma um fato sobre o passado. Daí a validação completa."""
    u = current_user()
    d = request.get_json(silent=True) or {}

    ref_tipo = (d.get("ref_tipo") or "livre").strip()
    ref_id = d.get("ref_id")
    titulo = d.get("titulo") or ""
    erro = core.validar_ref(ref_tipo, ref_id, titulo)
    if erro:
        return jsonify({"erro": erro}), 400

    inicio = core.parse_dt(d.get("inicio"))
    fim = core.parse_dt(d.get("fim"))
    erro = core.validar_intervalo(inicio, fim)
    if erro:
        return jsonify({"erro": erro}), 400
    if fim is None:
        return jsonify({"erro": "Informe o fim — lançamento manual não fica aberto"}), 400

    choque = core.conflita(u.id, inicio, fim)
    if choque:
        return jsonify({
            "erro": (f"Esse intervalo se sobrepõe a "
                     f"“{choque.rotulo()}” ({choque.inicio:%H:%M}–"
                     f"{choque.fim:%H:%M})" if choque.fim else
                     f"Esse intervalo colide com o cronômetro em curso"),
            "conflito": choque.to_dict(),
        }), 409

    ap = core.abrir(u, ref_tipo=ref_tipo, ref_id=ref_id, titulo=titulo,
                    nota=d.get("nota") or "", projeto_id=d.get("projeto_id"),
                    origem="manual", inicio=inicio, fim=fim)
    db.session.commit()

    log_action(u.email, "CREATE", entidade=f"Apontamento #{ap.id}",
               novo=f"{fmt_dur(ap.duracao_s)} em {ap.rotulo()}", ip=get_client_ip())
    _emit(EventType.TEMPO_APONTAMENTO_CREATED,
          {"entidade": "apontamento", "id": ap.id, "user_id": ap.user_id,
           "projeto_id": ap.projeto_id, "duracao_s": ap.duracao_s})
    return jsonify(ap.to_dict()), 201


@tempos_bp.route("/api/apontamentos/<int:ap_id>", methods=["GET"])
@so_membros
def detalhe(ap_id):
    ap, negado = _meu_ou_403(Apontamento.query.get_or_404(ap_id))
    if negado:
        return negado
    return jsonify(ap.to_dict())


@tempos_bp.route("/api/apontamentos/<int:ap_id>", methods=["PUT"])
@so_membros
def atualizar(ap_id):
    ap, negado = _meu_ou_403(Apontamento.query.get_or_404(ap_id))
    if negado:
        return negado
    u = current_user()
    d = request.get_json(silent=True) or {}
    antes = f"{fmt_dur(ap.duracao_s)} em {ap.rotulo()}"

    if "ref_tipo" in d or "ref_id" in d or "titulo" in d:
        ref_tipo = (d.get("ref_tipo") or ap.ref_tipo).strip()
        ref_id = d.get("ref_id", ap.ref_id)
        titulo = d.get("titulo", ap.titulo) or ""
        erro = core.validar_ref(ref_tipo, ref_id, titulo)
        if erro:
            return jsonify({"erro": erro}), 400
        alvo_mudou = (ref_tipo != ap.ref_tipo) or (ref_id != ap.ref_id)
        ap.ref_tipo = ref_tipo
        ap.ref_id = None if ref_tipo == "livre" else ref_id
        ap.titulo = titulo.strip()[:200]
        ap.titulo_chave = core.chave_titulo(titulo)
        # O rollup só é refeito quando o ALVO muda, ou quando o projeto vem
        # explícito no corpo. Recalcular incondicionalmente apagava o projeto de
        # uma tarefa livre a cada correção de título: para 'livre',
        # resolver_projeto devolve o projeto_id DO CORPO, que numa edição só de
        # título não vem — e o vínculo escolhido na criação virava None em
        # silêncio, levando as horas embora do relatório do projeto.
        if alvo_mudou or "projeto_id" in d:
            ap.projeto_id = core.resolver_projeto(ref_tipo, ref_id,
                                                  d.get("projeto_id"))

    if "nota" in d:
        ap.nota = (d.get("nota") or "").strip()

    # Horários. Na linha ABERTA só o início é editável ("liguei o cronômetro 20
    # minutos atrasado" é real); `fim` continua sendo trabalho exclusivo do
    # /parar, porque dois jeitos de sair do estado aberto é um a mais.
    novo_inicio, erro = _campo_dt(d, "inicio")
    if erro:
        return erro
    novo_fim, erro = _campo_dt(d, "fim")
    if erro:
        return erro

    if ap.fim is None:
        if novo_fim is not None:
            return jsonify({"erro": "Use Parar para encerrar o cronômetro"}), 400
        if novo_inicio:
            if novo_inicio > datetime.now():
                return jsonify({"erro": "Não dá para apontar tempo no futuro"}), 400
            ap.inicio = novo_inicio
    else:
        inicio = novo_inicio or ap.inicio
        fim = novo_fim or ap.fim
        erro = core.validar_intervalo(inicio, fim)
        if erro:
            return jsonify({"erro": erro}), 400
        choque = core.conflita(ap.user_id, inicio, fim, ignorar_id=ap.id)
        if choque:
            return jsonify({"erro": f"Esse intervalo se sobrepõe a “{choque.rotulo()}”",
                            "conflito": choque.to_dict()}), 409
        ap.inicio, ap.fim = inicio, fim
        ap.duracao_s = int((fim - inicio).total_seconds())

    # Gestor mexendo em hora alheia fica VISÍVEL. O MVP não tem aprovação; um
    # gestor reescrevendo horas em silêncio é aprovação com o consentimento
    # removido. Marcar quem editou mantém o poder honesto sem máquina de estado.
    ap.editado_por = u.email
    ap.editado_em = datetime.now()
    db.session.commit()

    log_action(u.email, "UPDATE", entidade=f"Apontamento #{ap.id}", antigo=antes,
               novo=f"{fmt_dur(ap.duracao_s)} em {ap.rotulo()}", ip=get_client_ip())
    _emit(EventType.TEMPO_APONTAMENTO_UPDATED,
          {"entidade": "apontamento", "id": ap.id, "user_id": ap.user_id,
           "projeto_id": ap.projeto_id, "duracao_s": ap.duracao_s})
    return jsonify(ap.to_dict())


@tempos_bp.route("/api/apontamentos/<int:ap_id>", methods=["DELETE"])
@so_membros
def excluir(ap_id):
    """Exclusão real, diferente de Composicao (arquiva) e Documento (soft delete).

    Aqueles são registros compartilhados cuja ausência confundiria outra pessoa.
    Um apontamento é pessoal, e um errado deve sair da folha em vez de virar linha
    `ativo=False` que todo SUM precisa lembrar de filtrar. O rastro fica no
    audit log.
    """
    ap, negado = _meu_ou_403(Apontamento.query.get_or_404(ap_id))
    if negado:
        return negado
    u = current_user()
    resumo = f"{fmt_dur(ap.duracao_s)} em {ap.rotulo()} ({ap.inicio:%d/%m/%Y %H:%M})"
    dono, apid, projeto = ap.user_id, ap.id, ap.projeto_id

    db.session.delete(ap)
    db.session.commit()

    log_action(u.email, "DELETE", entidade=f"Apontamento #{apid}", antigo=resumo,
               ip=get_client_ip())
    _emit(EventType.TEMPO_APONTAMENTO_DELETED,
          {"entidade": "apontamento", "id": apid, "user_id": dono,
           "projeto_id": projeto})
    return jsonify({"ok": True})


# ── RESUMOS ───────────────────────────────────────────────────────────────────

@tempos_bp.route("/api/resumo", methods=["GET"])
@so_membros
def resumo():
    q, erro = _lista_filtrada()
    if erro:
        return erro
    por = (request.args.get("por") or "dia").strip()
    linhas = q.limit(5000).all()
    # O mapa só é necessário no agrupamento por alvo; nos outros seria uma
    # consulta paga para nada. `agregar` computa sozinho se não receber.
    grupos = core.agregar(linhas, por,
                          refs=core.mapa_refs(linhas) if por == "alvo" else None)
    total = sum(core.segundos(a) for a in linhas)
    return jsonify({
        "por": por if por in core.AGRUPAMENTOS else "dia",
        "grupos": [dict(g, duracao=fmt_dur(g["duracao_s"])) for g in grupos],
        "total_s": total,
        "total": fmt_dur(total),
        "linhas": len(linhas),
    })


@tempos_bp.route("/api/por-ref", methods=["GET"])
@so_membros
def por_ref():
    """Total de um alvo — o número que as fichas de Missões/Equipamentos/Projetos
    exibem. Não é escopado por dono de propósito: "quanto o time já gastou neste
    equipamento" é dado de projeto, não de pessoa. Devolve só o agregado; nunca
    quem fez o quê."""
    ref_tipo = (request.args.get("ref_tipo") or "").strip()
    ref_id = (request.args.get("ref_id") or "").strip()
    if ref_tipo not in REF_TIPOS_TEMPO or not ref_id.isdigit():
        return jsonify({"erro": "informe ref_tipo e ref_id"}), 400

    rid = int(ref_id)
    do_alvo = (Apontamento.ref_tipo == ref_tipo, Apontamento.ref_id == rid)

    # Soma no BANCO, e não em Python. Esta rota é chamada a cada abertura de
    # ficha em Missões/Equipamentos/Projetos, e materializar o histórico inteiro
    # do alvo só para somá-lo cresce sem teto junto com a idade do sistema — um
    # equipamento com três anos de apontamento carregava três anos de linhas para
    # devolver um número.
    fechado_s, fechadas = (
        db.session.query(db.func.coalesce(db.func.sum(Apontamento.duracao_s), 0),
                         db.func.count(Apontamento.id))
        .filter(*do_alvo, Apontamento.fim.isnot(None))
        .one())

    # As abertas continuam em Python: `duracao_viva` conta o relógio correndo, e
    # isso o SQL não sabe fazer. A lista é curta por construção — no máximo uma
    # por pessoa, garantido por uq_tempo_em_curso.
    abertas = Apontamento.query.filter(*do_alvo, Apontamento.fim.is_(None)).all()
    total = int(fechado_s or 0) + sum(core.segundos(a) for a in abertas)

    pessoas = (db.session.query(db.func.count(db.distinct(Apontamento.user_id)))
               .filter(*do_alvo).scalar())

    return jsonify({"ref_tipo": ref_tipo, "ref_id": rid,
                    "total_s": total, "total": fmt_dur(total),
                    "linhas": int(fechadas or 0) + len(abertas),
                    "pessoas": int(pessoas or 0)})


# ── EXPORTS ───────────────────────────────────────────────────────────────────

_CAB = ["Data", "Pessoa", "Alvo", "Tipo do alvo", "Projeto", "Inicio", "Fim",
        "Duracao (h)", "Duracao", "Origem", "Encerrado por", "Nota"]


def _linha_export(ap, ref, projetos):
    return [
        ap.inicio.strftime("%d/%m/%Y") if ap.inicio else "",
        ap.usuario.nome if ap.usuario else "",
        ap.rotulo(ref),
        REF_TIPOS_LABELS.get(ap.ref_tipo, ap.ref_tipo),
        projetos.get(ap.projeto_id, ""),
        ap.inicio.strftime("%H:%M") if ap.inicio else "",
        ap.fim.strftime("%H:%M") if ap.fim else "",
        round(core.segundos(ap) / 3600, 2),
        fmt_dur(core.segundos(ap)),
        ap.origem or "",
        ap.encerrado_por or "",
        ap.nota or "",
    ]


def _dados_export():
    q, erro = _lista_filtrada()
    if erro:
        return None, None, None, erro
    linhas = q.limit(20000).all()
    refs = core.mapa_refs(linhas)
    ids = {a.projeto_id for a in linhas if a.projeto_id}
    projetos = ({p.id: p.nome for p in Projeto.query.filter(Projeto.id.in_(ids)).all()}
                if ids else {})
    return linhas, refs, projetos, None


@tempos_bp.route("/api/export/apontamentos.csv", methods=["GET"])
@so_membros
def exportar_csv():
    linhas, refs, projetos, erro = _dados_export()
    if erro:
        return erro
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";")
    w.writerow(_CAB)
    for ap in linhas:
        w.writerow(_linha_export(ap, refs.get(ap.id), projetos))
    out = io.BytesIO(buf.getvalue().encode("utf-8-sig"))
    return send_file(out, mimetype="text/csv", as_attachment=True,
                     download_name=f"tempos_{datetime.now():%Y%m%d}.csv")


@tempos_bp.route("/api/export/tempos.xlsx", methods=["GET"])
@so_membros
def exportar_xlsx():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    linhas, refs, projetos, erro = _dados_export()
    if erro:
        return erro

    wb = Workbook()
    cab_fonte = Font(bold=True, color="FFFFFF")
    cab_fundo = PatternFill("solid", fgColor="1F4E5F")

    def nova_aba(titulo, cabecalho, larguras, primeira=False):
        ws = wb.active if primeira else wb.create_sheet()
        ws.title = titulo
        ws.append(cabecalho)
        for cel in ws[1]:
            cel.font, cel.fill = cab_fonte, cab_fundo
        for letra, larg in larguras.items():
            ws.column_dimensions[letra].width = larg
        ws.freeze_panes = "A2"
        return ws

    ws = nova_aba("Apontamentos", _CAB,
                  {"A": 12, "B": 22, "C": 36, "D": 18, "E": 26, "L": 40},
                  primeira=True)
    for ap in linhas:
        ws.append(_linha_export(ap, refs.get(ap.id), projetos))

    wr = nova_aba("Por projeto", ["Projeto", "Horas", "Duracao", "Lancamentos"],
                  {"A": 34, "C": 14})
    for g in core.agregar(linhas, "projeto"):
        wr.append([g["rotulo"], round(g["duracao_s"] / 3600, 2),
                   fmt_dur(g["duracao_s"]), g["linhas"]])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(
        buf, as_attachment=True,
        download_name=f"tempos_{datetime.now():%Y%m%d}.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
